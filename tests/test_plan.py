#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from pathlib import Path
import os
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from generate_catalog_ppt import (  # noqa: E402
    build_deck_plan,
    detect_language,
    extract_page_count,
    extract_product_categories,
    extract_requested_sections,
    generate,
    sanitize_filename,
    category_present,
    expected_explicit_categories_for_plan,
)
from pptx import Presentation  # noqa: E402


SELLERS_UNION_PROMPT = (
    "Generate a 10-page festive party theme product catalog PPT for SELLERS UNION. "
    "Page 1: Company introduction for SELLERS UNION - a premier party supplies and festive decorations supplier. "
    "Pages 2-10: Each page features one product category with party-related images and English descriptions. "
    "Product categories: Balloons and Balloon Sets, Rain Curtains, Candles, Bunting Garlands, Hats, Party Blowers, "
    "Disposable Party Tableware, COS Costumes, and other party decorations. All text in English. "
    "High-quality commercial product catalog style."
)


def test_detailed_english_plan():
    plan = build_deck_plan(SELLERS_UNION_PROMPT)
    assert plan.language == "en"
    assert plan.page_count == 10
    assert len(plan.slides) == 10
    assert plan.slides[0].layout == "company_intro"
    assert plan.slides[0].title == "SELLERS UNION"
    assert plan.slides[1].category == "Balloons and Balloon Sets"
    assert plan.slides[-1].category == "Other Party Decorations"
    assert len(plan.image_briefs) == 10


def test_extract_categories_does_not_split_balloon_and_sets():
    cats = extract_product_categories(SELLERS_UNION_PROMPT, "en")
    assert cats[0] == "Balloons and Balloon Sets"
    assert len(cats) == 9


def test_page_count_patterns():
    assert extract_page_count("Generate a 10-page catalog") == 10
    assert extract_page_count("Generate 12 pages catalog") == 12
    assert extract_page_count("生成一个厨房餐具相关的3页的商品目录册PPT") == 3
    assert extract_page_count("做一份厨房餐具三页PPT") == 3


def test_language_detection():
    assert detect_language(SELLERS_UNION_PROMPT) == "en"
    assert detect_language("生成一个厨房餐具相关的3页的商品目录册PPT") == "zh"


def test_filename_keeps_normal_english_letters():
    name = sanitize_filename("Generate a 10-page festive party theme product catalog PPT for SELLERS UNION")
    assert "Generate" in name
    assert "party" in name


def test_placeholder_generation_exact_10_pages_and_english_text():
    with tempfile.TemporaryDirectory() as td:
        os.environ["PPT_REQUIRE_IMAGES"] = "1"
        result = generate(SELLERS_UNION_PROMPT, root=td, image_mode="placeholder")
        assert result.ok, result.error
        prs = Presentation(result.pptx_path)
        assert len(prs.slides) == 10
        text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
        assert "商品" not in text
        assert "Balloons and Balloon Sets" in text
        assert "Other Party Decorations" in text
        assert result.validation["ok"] is True


def test_categories_can_fill_all_pages_without_forced_cover():
    prompt = "Generate a 3-page minimalist stationery product catalog PPT. Product categories: Notebooks, Pens, Erasers. All text in English."
    plan = build_deck_plan(prompt)
    assert plan.page_count == 3
    assert [s.layout for s in plan.slides] == ["category_showcase", "category_showcase", "category_showcase"]
    assert [s.title for s in plan.slides] == ["Notebooks", "Pens", "Erasers"]
    assert plan.style_family == "minimal"


def test_unrequested_procurement_fields_are_not_forced():
    plan = build_deck_plan(SELLERS_UNION_PROMPT)
    section_titles = [section["title"] for slide in plan.slides for section in slide.sections]
    assert "Suggested SKUs" not in section_titles
    assert "Packaging & Sourcing Notes" not in section_titles
    assert "MOQ Note" not in section_titles
    assert "Selling Points" not in section_titles


def test_negative_field_requirements_are_respected():
    prompt = "生成5页文具目录册PPT，类目：笔记本，笔，橡皮，尺子，修正带。只要产品方向，不要价格MOQ箱规。"
    assert extract_requested_sections(prompt) == ["positioning"]
    cats = extract_product_categories(prompt, "zh")
    assert cats == ["笔记本", "笔", "橡皮", "尺子", "修正带"]
    plan = build_deck_plan(prompt)
    assert len(plan.slides) == 5
    assert all([section["title"] for section in slide.sections] == ["产品定位"] for slide in plan.slides)


def test_explicit_procurement_fields_are_added_only_when_requested():
    prompt = "生成3页高端文具目录册PPT，类目：笔记本，笔。需要卖点、包装、MOQ、认证。"
    plan = build_deck_plan(prompt)
    titles = [section["title"] for slide in plan.slides for section in slide.sections]
    assert "MOQ 备注" in titles
    assert "认证信息" in titles
    assert "包装方式" in titles
    assert plan.style_family == "luxury"


def test_category_match_is_robust_for_wrapped_long_titles():
    assert category_present("Balloons and Balloon Sets", ["Balloons\n& Balloon Sets"])
    assert category_present("Disposable Party Tableware", ["Disposable\nParty\nTableware"])


def test_expected_categories_uses_actual_category_slide_capacity():
    plan = build_deck_plan(SELLERS_UNION_PROMPT)
    expected = expected_explicit_categories_for_plan(plan)
    assert expected == plan.explicit_categories
    assert "Balloons and Balloon Sets" in expected
    assert "Disposable Party Tableware" in expected


if __name__ == "__main__":
    test_detailed_english_plan()
    test_extract_categories_does_not_split_balloon_and_sets()
    test_page_count_patterns()
    test_language_detection()
    test_filename_keeps_normal_english_letters()
    test_placeholder_generation_exact_10_pages_and_english_text()
    test_categories_can_fill_all_pages_without_forced_cover()
    test_unrequested_procurement_fields_are_not_forced()
    test_negative_field_requirements_are_respected()
    test_explicit_procurement_fields_are_added_only_when_requested()
    test_category_match_is_robust_for_wrapped_long_titles()
    test_expected_categories_uses_actual_category_slide_capacity()
    print("ok")
