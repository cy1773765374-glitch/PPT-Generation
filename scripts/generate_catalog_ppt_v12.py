#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
OpenClaw PPT Generation v1.2

核心变化：
- 飞书回复只返回 Windows/Samba 本地路径，不再回复服务器路径/耗时/页数。
- 文件名时间固定使用 Asia/Shanghai，格式为：YYYY-MM-DD-HH时MM分-用户询问的问题。
- 生产默认强制调用 MiniMax image-01 生图，并将图片插入 PPT；如果 MiniMax 不可用则失败，不生成纯文字 PPT。
- PPT 文件仍直接放在任务目录一级，不使用 project/exports/。
"""

from __future__ import annotations

import argparse
import base64
import binascii
import json
import mimetypes
import os
import re
import sys
import time
import urllib.parse
from dataclasses import dataclass, asdict, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from zoneinfo import ZoneInfo

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover
    load_dotenv = None

try:
    import requests
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"requests 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)

try:
    from PIL import Image, ImageDraw, ImageFont
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"Pillow 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"python-pptx 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)


CN_NUM_MAP = {
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}

INVALID_FILENAME_CHARS = r'<>:"/\\|?*\n\r\t'
DEFAULT_TZ = "Asia/Shanghai"


@dataclass
class CatalogSpec:
    prompt: str
    category: str
    page_count: int
    doc_type: str
    sender_name: str = ""
    sender_open_id: str = ""


@dataclass
class ImageAsset:
    key: str
    title: str
    prompt: str
    path: str
    source: str
    aspect_ratio: str


@dataclass
class GenerationResult:
    ok: bool
    prompt: str
    category: str
    page_count: int
    run_name: str
    run_dir: str
    pptx_path: str
    windows_path: str
    elapsed_seconds: float
    image_mode: str = ""
    image_count: int = 0
    image_assets: List[Dict[str, Any]] = field(default_factory=list)
    reply_text: str = ""
    error: str = ""


def env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if not value:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or value == "":
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def now_in_configured_timezone() -> datetime:
    tz_name = os.getenv("PPT_TIMEZONE", DEFAULT_TZ).strip() or DEFAULT_TZ
    try:
        tz = ZoneInfo(tz_name)
    except Exception:
        tz = ZoneInfo(DEFAULT_TZ)
    return datetime.now(tz)


def parse_cn_number(text: str) -> Optional[int]:
    text = text.strip()
    if not text:
        return None
    if text.isdigit():
        return int(text)
    if text in CN_NUM_MAP:
        return CN_NUM_MAP[text]
    if text.startswith("十"):
        if len(text) == 1:
            return 10
        return 10 + CN_NUM_MAP.get(text[1:], 0)
    if "十" in text:
        left, _, right = text.partition("十")
        left_num = CN_NUM_MAP.get(left, 1)
        right_num = CN_NUM_MAP.get(right, 0) if right else 0
        return left_num * 10 + right_num
    return None


def extract_page_count(prompt: str, default: int, max_page_count: int) -> int:
    patterns = [
        r"(\d{1,3})\s*[页頁]",
        r"([一二两三四五六七八九十]{1,3})\s*[页頁]",
        r"(\d{1,3})\s*page",
        r"(\d{1,3})\s*slides?",
    ]
    for pattern in patterns:
        match = re.search(pattern, prompt, flags=re.IGNORECASE)
        if match:
            value = parse_cn_number(match.group(1))
            if value:
                return max(1, min(value, max_page_count))
    return max(1, min(default, max_page_count))


def clean_prompt(prompt: str) -> str:
    prompt = re.sub(r"\s+", " ", prompt.strip())
    # 去掉飞书引用/回复前缀，避免把“回复 陈玉:”纳入目录名。
    prompt = re.sub(r"^回复\s*[^:：]{1,30}[:：]\s*", "", prompt)
    # 如果上游错误地把生成结果也喂回来，只保留第一行真实需求。
    for marker in ["已完成，PPT 文件", "服务器路径", "已完成 ·", "PPT 文件："]:
        if marker in prompt:
            prompt = prompt.split(marker, 1)[0].strip()
    return prompt.strip()


def remove_task_words(text: str) -> str:
    words = [
        "生成", "做", "制作", "设计", "帮我", "帮我做", "帮我生成",
        "一个", "一份", "相关", "商品目录册", "产品目录册", "目录册",
        "PPT", "ppt", "幻灯片", "的", "页", "页面",
    ]
    result = text.strip()
    for word in words:
        result = result.replace(word, "")
    result = re.sub(r"\d+", "", result)
    result = re.sub(r"[一二两三四五六七八九十]+", "", result)
    return result.strip(" -_，。,.：:；;")


def extract_category(prompt: str) -> str:
    text = clean_prompt(prompt)
    match = re.search(r"(?:生成|做|制作|出|帮我做|帮我生成|设计)?(?:一个|一份)?(.{2,30}?)(?:相关|类|系列)的?", text)
    if match:
        category = remove_task_words(match.group(1))
        if category:
            return category

    match = re.search(r"(?:生成|做|制作|设计|帮我生成|帮我做)(?:一个|一份)?(.{2,30}?)(?:的)?(?:\d+|[一二两三四五六七八九十]+)?页", text)
    if match:
        category = remove_task_words(match.group(1))
        if category:
            return category

    match = re.search(r"(.{2,30}?)(?:商品目录册|产品目录册|目录册|PPT|ppt)", text)
    if match:
        category = remove_task_words(match.group(1))
        if category:
            return category

    return "综合商品"


def sanitize_filename(text: str, max_chars: int = 56) -> str:
    text = clean_prompt(text)
    for ch in INVALID_FILENAME_CHARS:
        text = text.replace(ch, "")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，,。.!！?？：:；;]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-._ ")
    if not text:
        text = "商品目录册"
    if len(text) > max_chars:
        text = text[:max_chars].rstrip("-._ ")
    return text


def make_run_paths(prompt: str, root: Path, now: datetime) -> Dict[str, Path | str]:
    max_chars = env_int("PPT_FILENAME_MAX_CHARS", 56)
    # 用户要求“2026-05-25-几时几分-用户询问的问题”；这里用 16时35分，避免 0935 的歧义。
    time_prefix = now.strftime("%Y-%m-%d-%H时%M分")
    slug = sanitize_filename(prompt, max_chars=max_chars)
    base_name = f"{time_prefix}-{slug}"

    auto_dedup = env_bool("PPT_AUTO_DEDUP", True)
    run_name = base_name
    run_dir = root / run_name
    if auto_dedup:
        i = 1
        while run_dir.exists():
            run_name = f"{base_name}-{i:02d}"
            run_dir = root / run_name
            i += 1

    pptx_path = run_dir / f"{run_name}.pptx"
    return {"run_name": run_name, "run_dir": run_dir, "pptx_path": pptx_path}


def windows_path_for(pptx_path: Path, root: Path) -> str:
    prefix = os.getenv("PPT_WINDOWS_SHARE_PREFIX", "Z:\\yaq\\ppt\\catalog").rstrip("\\/")
    try:
        rel = pptx_path.relative_to(root)
    except ValueError:
        rel = pptx_path.name
    rel_win = str(rel).replace("/", "\\")
    return f"{prefix}\\{rel_win}"


def product_ideas(category: str) -> List[Dict[str, Any]]:
    category_key = category.replace(" ", "")
    if any(k in category_key for k in ["厨房餐具", "餐具", "厨具", "厨房"]):
        return [
            {
                "name": "不锈钢刀叉勺套装",
                "position": "家庭日用 / 礼品套装",
                "selling_points": ["食品级不锈钢", "镜面抛光", "多规格组合"],
                "spec": "16件/24件/30件可选；独立彩盒包装",
                "visual": "premium stainless steel fork spoon knife cutlery set, neatly arranged, modern kitchen product photography",
            },
            {
                "name": "陶瓷碗盘组合",
                "position": "餐桌场景 / 家居零售",
                "selling_points": ["高温烧制", "釉面细腻", "适配多风格餐桌"],
                "spec": "碗、盘、汤碗、调味碟组合；颜色可定制",
                "visual": "elegant ceramic bowls and plates dinnerware set on a warm dining table, clean catalog photography",
            },
            {
                "name": "硅胶厨具铲勺套装",
                "position": "厨房烹饪 / 防刮锅具",
                "selling_points": ["耐高温硅胶", "不伤锅涂层", "易清洗收纳"],
                "spec": "5件/8件/12件套；可加收纳桶",
                "visual": "food grade silicone cooking utensils set with spatulas and ladles, organized in a holder, bright kitchen scene",
            },
            {
                "name": "厨房餐具收纳架",
                "position": "厨房收纳 / 台面整理",
                "selling_points": ["分区收纳", "沥水设计", "节省台面空间"],
                "spec": "单层/双层；金属/塑料材质可选",
                "visual": "compact kitchen cutlery organizer rack with draining design, utensils arranged neatly, product catalog style",
            },
        ]
    return [
        {
            "name": f"{category}核心款 A",
            "position": "主推款 / 常规渠道",
            "selling_points": ["高频使用", "稳定供应", "支持定制"],
            "spec": "尺寸、颜色、包装、MOQ 建议确认",
            "visual": f"{category} hero product, commercial product catalog photography, clean background, premium lighting",
        },
        {
            "name": f"{category}升级款 B",
            "position": "中高端款 / 礼品渠道",
            "selling_points": ["质感升级", "外观差异化", "适合套装销售"],
            "spec": "材质、工艺、认证、价格建议确认",
            "visual": f"upgraded {category} product set, premium e-commerce product photography, clean background",
        },
        {
            "name": f"{category}组合款 C",
            "position": "套装款 / 批量采购",
            "selling_points": ["组合灵活", "性价比高", "便于陈列"],
            "spec": "组合数量、箱规、交期建议确认",
            "visual": f"{category} product bundle set, neat arrangement, professional catalog photography",
        },
        {
            "name": f"{category}展示款 D",
            "position": "形象款 / 目录册展示",
            "selling_points": ["视觉统一", "适合封面", "支持品牌化"],
            "spec": "颜色、LOGO、包装方案建议确认",
            "visual": f"{category} product display, lifestyle scene, commercial catalog photography",
        },
    ]


def read_openclaw_minimax_key() -> str:
    """尽量复用 OpenClaw 配置中的 MiniMax API Key，不把密钥写入本包。"""
    config_path = Path(os.getenv("OPENCLAW_CONFIG", "~/.openclaw/openclaw.json")).expanduser()
    if not config_path.exists():
        return ""
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return ""

    key_names = {"apikey", "api_key", "key", "token", "authorization"}

    def walk(obj: Any, parent_keys: List[str]) -> Iterable[str]:
        if isinstance(obj, dict):
            for k, v in obj.items():
                lk = str(k).lower()
                new_parent = parent_keys + [lk]
                if isinstance(v, str) and lk in key_names and any("minimax" in p or "minimaxi" in p for p in parent_keys):
                    yield v
                else:
                    yield from walk(v, new_parent)
        elif isinstance(obj, list):
            for item in obj:
                yield from walk(item, parent_keys)

    for candidate in walk(data, []):
        candidate = candidate.strip()
        if candidate and not candidate.startswith("${") and len(candidate) > 8:
            if candidate.lower().startswith("bearer "):
                candidate = candidate.split(None, 1)[1].strip()
            return candidate
    return ""


def get_minimax_api_key() -> str:
    for name in ["MINIMAX_API_KEY", "MINIMAX_TOKEN", "MINIMAX_API_TOKEN"]:
        value = os.getenv(name, "").strip()
        if value:
            return value.removeprefix("Bearer ").strip()
    key_file = os.getenv("MINIMAX_API_KEY_FILE", "").strip()
    if key_file:
        path = Path(key_file).expanduser()
        if path.exists():
            return path.read_text(encoding="utf-8").strip().removeprefix("Bearer ").strip()
    return read_openclaw_minimax_key()


def collect_strings(obj: Any) -> Iterable[str]:
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for value in obj.values():
            yield from collect_strings(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from collect_strings(value)


def looks_like_base64_image(text: str) -> bool:
    if text.startswith("data:image/"):
        return True
    if len(text) < 1000:
        return False
    compact = text.strip().replace("\n", "")
    if len(compact) % 4 not in {0, 2, 3}:
        return False
    return re.fullmatch(r"[A-Za-z0-9+/=]+", compact) is not None


def decode_base64_image(text: str) -> bytes:
    if text.startswith("data:image/"):
        _, _, text = text.partition(",")
    compact = text.strip().replace("\n", "")
    missing = len(compact) % 4
    if missing:
        compact += "=" * (4 - missing)
    return base64.b64decode(compact, validate=False)


class MiniMaxImageClient:
    def __init__(self) -> None:
        self.api_key = get_minimax_api_key()
        self.api_url = os.getenv("MINIMAX_IMAGE_API_URL", "https://api.minimax.io/v1/image_generation").strip()
        self.model = os.getenv("MINIMAX_IMAGE_MODEL", "image-01").strip()
        self.timeout = env_int("MINIMAX_IMAGE_TIMEOUT", 180)
        self.response_format = os.getenv("MINIMAX_IMAGE_RESPONSE_FORMAT", "base64").strip().lower()
        self.prompt_optimizer = env_bool("MINIMAX_PROMPT_OPTIMIZER", True)
        self.verify_ssl = env_bool("MINIMAX_VERIFY_SSL", True)

    def generate_to_file(self, prompt: str, out_path: Path, aspect_ratio: str = "16:9") -> None:
        if not self.api_key:
            raise RuntimeError("MINIMAX_API_KEY 未配置，且未能从 ~/.openclaw/openclaw.json 自动读取到 MiniMax 密钥；为避免生成纯文字 PPT，本次已停止。")

        payload: Dict[str, Any] = {
            "model": self.model,
            "prompt": prompt[:1500],
            "aspect_ratio": aspect_ratio,
            "response_format": self.response_format,
            "n": 1,
            "prompt_optimizer": self.prompt_optimizer,
        }
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json",
        }
        resp = requests.post(self.api_url, headers=headers, json=payload, timeout=self.timeout, verify=self.verify_ssl)
        if resp.status_code >= 400:
            raise RuntimeError(f"MiniMax 生图失败 HTTP {resp.status_code}: {resp.text[:600]}")
        try:
            data = resp.json()
        except Exception as exc:
            raise RuntimeError(f"MiniMax 生图返回不是 JSON：{exc}; body={resp.text[:300]}") from exc

        base_resp = data.get("base_resp") if isinstance(data, dict) else None
        if isinstance(base_resp, dict) and int(base_resp.get("status_code", 0) or 0) != 0:
            raise RuntimeError(f"MiniMax 生图失败：{base_resp.get('status_msg') or base_resp}")

        out_path.parent.mkdir(parents=True, exist_ok=True)

        image_urls = []
        if isinstance(data, dict):
            image_urls = data.get("data", {}).get("image_urls", []) if isinstance(data.get("data"), dict) else []
        if image_urls:
            self._download_image(str(image_urls[0]), out_path)
            return

        for item in collect_strings(data):
            if item.startswith("http://") or item.startswith("https://"):
                self._download_image(item, out_path)
                return
            if looks_like_base64_image(item):
                raw = decode_base64_image(item)
                if len(raw) > 1000:
                    out_path.write_bytes(raw)
                    self._normalize_image(out_path)
                    return

        raise RuntimeError(f"MiniMax 生图响应中未找到 image_urls 或 base64 图片字段：{json.dumps(data, ensure_ascii=False)[:800]}")

    def _download_image(self, url: str, out_path: Path) -> None:
        resp = requests.get(url, timeout=self.timeout, verify=self.verify_ssl)
        if resp.status_code >= 400:
            raise RuntimeError(f"MiniMax 图片下载失败 HTTP {resp.status_code}: {url}")
        out_path.write_bytes(resp.content)
        self._normalize_image(out_path)

    def _normalize_image(self, path: Path) -> None:
        try:
            with Image.open(path) as im:
                im = im.convert("RGB")
                im.save(path, format="PNG")
        except Exception as exc:
            raise RuntimeError(f"MiniMax 图片文件无法被 Pillow 读取：{path}；{exc}") from exc


def make_catalog_image_prompts(spec: CatalogSpec, products: List[Dict[str, Any]]) -> List[Dict[str, str]]:
    base_style = (
        "commercial product catalog photography, realistic, premium lighting, clean warm white background, "
        "modern e-commerce catalog, no text, no watermark, no logo, high detail"
    )
    prompts = [
        {
            "key": "cover_hero",
            "title": f"{spec.category}封面场景图",
            "aspect_ratio": "16:9",
            "prompt": f"A premium product catalog hero image for {spec.category}: assorted kitchen tableware and utensils arranged beautifully on a modern kitchen counter, {base_style}",
        }
    ]
    for i, product in enumerate(products[:4], start=1):
        prompts.append(
            {
                "key": f"product_{i}",
                "title": product["name"],
                "aspect_ratio": "1:1",
                "prompt": f"{product.get('visual', product['name'])}, product only, centered composition, {base_style}",
            }
        )
    return prompts


def generate_placeholder_image(title: str, out_path: Path, aspect_ratio: str) -> None:
    """仅用于离线测试；生产不要启用 PPT_IMAGE_MODE=placeholder。"""
    if aspect_ratio == "16:9":
        size = (1280, 720)
    else:
        size = (1024, 1024)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", size, (246, 248, 250))
    draw = ImageDraw.Draw(im)
    draw.rounded_rectangle([40, 40, size[0] - 40, size[1] - 40], radius=36, fill=(255, 255, 255), outline=(203, 213, 225), width=4)
    draw.ellipse([size[0] * 0.12, size[1] * 0.18, size[0] * 0.42, size[1] * 0.58], fill=(219, 234, 254))
    draw.rounded_rectangle([size[0] * 0.48, size[1] * 0.22, size[0] * 0.83, size[1] * 0.68], radius=28, fill=(224, 242, 254))
    draw.line([size[0] * 0.18, size[1] * 0.72, size[0] * 0.82, size[1] * 0.72], fill=(37, 99, 235), width=8)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 36)
    except Exception:
        font = ImageFont.load_default()
    text = f"OFFLINE PLACEHOLDER\n{title}"
    draw.multiline_text((70, size[1] - 165), text, fill=(71, 85, 105), font=font, spacing=10)
    im.save(out_path, format="PNG")


def prepare_images(spec: CatalogSpec, products: List[Dict[str, Any]], run_dir: Path, image_mode: str) -> List[ImageAsset]:
    image_mode = image_mode.lower().strip()
    if image_mode not in {"minimax", "placeholder", "off"}:
        raise ValueError("PPT_IMAGE_MODE 只能是 minimax / placeholder / off")
    if image_mode == "off":
        if env_bool("PPT_REQUIRE_IMAGES", True):
            raise RuntimeError("PPT_IMAGE_MODE=off 但 PPT_REQUIRE_IMAGES=1；为避免纯文字 PPT，本次已停止。")
        return []

    prompts = make_catalog_image_prompts(spec, products)
    assets: List[ImageAsset] = []
    asset_dir = run_dir / "assets" / image_mode
    client = MiniMaxImageClient() if image_mode == "minimax" else None

    for item in prompts:
        out_path = asset_dir / f"{item['key']}.png"
        if image_mode == "minimax":
            assert client is not None
            client.generate_to_file(item["prompt"], out_path, item["aspect_ratio"])
        else:
            generate_placeholder_image(item["title"], out_path, item["aspect_ratio"])
        assets.append(ImageAsset(key=item["key"], title=item["title"], prompt=item["prompt"], path=str(out_path), source=image_mode, aspect_ratio=item["aspect_ratio"]))

    if env_bool("PPT_REQUIRE_IMAGES", True) and len(assets) == 0:
        raise RuntimeError("未生成任何图片；为避免纯文字 PPT，本次已停止。")
    return assets


def asset_map(assets: List[ImageAsset]) -> Dict[str, ImageAsset]:
    return {asset.key: asset for asset in assets}


def set_slide_background(slide, rgb=(248, 250, 252)) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def add_textbox(slide, left, top, width, height, text, font_size=20, bold=False, color=(31, 41, 55), align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()
    lines = str(text).split("\n")
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            run.font.name = "Microsoft YaHei"
        if not p.runs:
            run = p.add_run()
            run.text = line
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            run.font.name = "Microsoft YaHei"
    return box


def add_title(slide, title: str, subtitle: str = "") -> None:
    add_textbox(slide, Inches(0.65), Inches(0.35), Inches(12.0), Inches(0.55), title, 27, True, (15, 23, 42))
    if subtitle:
        add_textbox(slide, Inches(0.68), Inches(0.9), Inches(11.8), Inches(0.35), subtitle, 11, False, (100, 116, 139))
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.22), Inches(12.0), Inches(0.035))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(37, 99, 235)
    line.line.fill.background()


def add_image(slide, asset: Optional[ImageAsset], left, top, width, height, label: str = "") -> None:
    if asset and Path(asset.path).exists():
        try:
            pic = slide.shapes.add_picture(asset.path, left, top, width=width, height=height)
            pic.line.color.rgb = RGBColor(226, 232, 240)
            return
        except Exception:
            pass
    # 生产上通常不会走到这里；兜底只是避免 python-pptx 崩溃。
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(226, 232, 240)
    shape.line.color.rgb = RGBColor(203, 213, 225)
    add_textbox(slide, left + Inches(0.18), top + height / 2 - Inches(0.18), width - Inches(0.36), Inches(0.5), label or "图片生成失败", 11, True, (100, 116, 139), PP_ALIGN.CENTER)


def add_card(slide, left, top, width, height, title: str, body_lines: List[str], accent=(37, 99, 235), font_size=10.2) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(226, 232, 240)
    shape.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()

    add_textbox(slide, left + Inches(0.22), top + Inches(0.14), width - Inches(0.35), Inches(0.28), title, 13.5, True, (15, 23, 42))
    box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.5), width - Inches(0.38), height - Inches(0.58))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, line in enumerate(body_lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(font_size)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = RGBColor(71, 85, 105)
        p.space_after = Pt(2)


def add_footer(slide, text: str) -> None:
    add_textbox(slide, Inches(0.68), Inches(7.06), Inches(12.0), Inches(0.22), text, 8.5, False, (148, 163, 184))


def create_cover(prs: Presentation, spec: CatalogSpec, run_name: str, assets: Dict[str, ImageAsset], now: datetime) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, (241, 245, 249))

    add_textbox(slide, Inches(0.78), Inches(0.72), Inches(5.4), Inches(0.85), f"{spec.category}\n商品目录册", 34, True, (15, 23, 42))
    add_textbox(slide, Inches(0.82), Inches(2.02), Inches(5.1), Inches(0.35), "Product Catalog · AI Generated", 15, False, (37, 99, 235))
    add_textbox(slide, Inches(0.82), Inches(2.62), Inches(5.15), Inches(0.78), f"输入需求：{spec.prompt}", 13.5, False, (51, 65, 85))
    add_card(slide, Inches(0.82), Inches(4.15), Inches(5.25), Inches(1.35), "目录册信息", [f"页数：{spec.page_count} 页", f"生成时间：{now.strftime('%Y-%m-%d %H:%M')}（Asia/Shanghai）", "版本：OpenClaw PPT Generation v1.2"], (37, 99, 235), 9.5)
    add_textbox(slide, Inches(0.82), Inches(6.12), Inches(5.2), Inches(0.4), "图片由 MiniMax image-01 生成并插入 PPT", 10.5, False, (100, 116, 139))

    img_bg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.45), Inches(0.72), Inches(6.25), Inches(5.85))
    img_bg.fill.solid()
    img_bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    img_bg.line.color.rgb = RGBColor(226, 232, 240)
    add_image(slide, assets.get("cover_hero"), Inches(6.68), Inches(0.95), Inches(5.8), Inches(3.28), "MiniMax 封面图")

    product_names = " / ".join([p["name"] for p in product_ideas(spec.category)[:4]])
    add_textbox(slide, Inches(6.75), Inches(4.55), Inches(5.65), Inches(0.75), product_names, 16, True, (15, 23, 42))
    add_textbox(slide, Inches(6.75), Inches(5.5), Inches(5.65), Inches(0.6), "面向业务初稿：可继续补充实物图、报价、MOQ、认证、箱规后生成正式客户版。", 11, False, (71, 85, 105))
    add_footer(slide, run_name)


def create_product_matrix(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int, assets: Dict[str, ImageAsset]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "商品系列矩阵", f"围绕“{spec.category}”生成的第一版目录册结构，图片由 MiniMax 生成")

    positions = [
        (Inches(0.75), Inches(1.5)),
        (Inches(6.85), Inches(1.5)),
        (Inches(0.75), Inches(4.15)),
        (Inches(6.85), Inches(4.15)),
    ]
    for idx, product in enumerate(products[:4]):
        left, top = positions[idx]
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.75), Inches(2.28))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = RGBColor(226, 232, 240)
        add_image(slide, assets.get(f"product_{idx+1}"), left + Inches(0.16), top + Inches(0.18), Inches(1.72), Inches(1.72), product["name"])
        add_textbox(slide, left + Inches(2.05), top + Inches(0.18), Inches(3.42), Inches(0.3), product["name"], 13, True, (15, 23, 42))
        lines = [
            f"定位：{product['position']}",
            f"卖点：{' / '.join(product['selling_points'])}",
            f"规格：{product['spec']}",
        ]
        box = slide.shapes.add_textbox(left + Inches(2.05), top + Inches(0.62), Inches(3.42), Inches(1.28))
        tf = box.text_frame
        tf.word_wrap = True
        tf.clear()
        for line_idx, line in enumerate(lines):
            p = tf.paragraphs[0] if line_idx == 0 else tf.add_paragraph()
            p.text = f"• {line}"
            p.font.size = Pt(9.2)
            p.font.name = "Microsoft YaHei"
            p.font.color.rgb = RGBColor(71, 85, 105)
            p.space_after = Pt(1)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), Inches(2.28))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor(37, 99, 235) if idx % 2 == 0 else RGBColor(14, 165, 233)
        accent.line.fill.background()
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_specs_page(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int, assets: Dict[str, ImageAsset]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "规格建议与采购沟通信息", "便于业务人员继续补充真实商品参数")

    add_image(slide, assets.get("cover_hero"), Inches(0.8), Inches(1.45), Inches(3.85), Inches(2.18), "MiniMax 场景图")
    add_card(slide, Inches(0.8), Inches(3.95), Inches(3.85), Inches(2.55), "沟通建议", ["补充真实图片、价格、材质、MOQ", "确认包装、箱规、认证与交期", "可继续输入品牌信息生成正式版"], (16, 185, 129), 9.5)

    headers = ["项目", "建议内容", "备注"]
    rows = [
        ["商品系列", " / ".join([p["name"] for p in products[:3]]), "按实际 SKU 调整"],
        ["材质", "不锈钢 / 陶瓷 / 食品级硅胶 / PP 等", "按产品确认"],
        ["包装", "彩盒、礼盒、裸装、外箱", "补充包装尺寸"],
        ["定制", "LOGO、颜色、组合数量、说明书", "确认 MOQ 与打样周期"],
        ["贸易信息", "价格、箱规、交期、认证", "业务人工补充"],
    ]

    table = slide.shapes.add_table(len(rows) + 1, 3, Inches(5.0), Inches(1.45), Inches(7.65), Inches(4.95)).table
    table.columns[0].width = Inches(1.25)
    table.columns[1].width = Inches(4.55)
    table.columns[2].width = Inches(1.85)

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(37, 99, 235)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(10.5)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.name = "Microsoft YaHei"

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9.2)
                p.font.color.rgb = RGBColor(51, 65, 85)
                p.font.name = "Microsoft YaHei"

    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_product_detail(prs: Presentation, spec: CatalogSpec, product: Dict[str, Any], page_no: int, asset: Optional[ImageAsset]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, product["name"], f"{spec.category} · 单品详情页")

    add_image(slide, asset, Inches(0.85), Inches(1.48), Inches(4.85), Inches(4.85), product["name"])
    add_card(slide, Inches(6.1), Inches(1.48), Inches(6.35), Inches(1.45), "产品定位", [product["position"], "适合用于目录册主推款展示"], (37, 99, 235))
    add_card(slide, Inches(6.1), Inches(3.22), Inches(6.35), Inches(1.45), "核心卖点", product["selling_points"], (14, 165, 233))
    add_card(slide, Inches(6.1), Inches(4.96), Inches(6.35), Inches(1.1), "规格建议", [product["spec"]], (16, 185, 129))
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_selling_points(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int, assets: Dict[str, ImageAsset]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "核心卖点与采购价值", "用于销售沟通、客户初筛、目录册初版展示")

    add_image(slide, assets.get("cover_hero"), Inches(0.8), Inches(1.48), Inches(4.2), Inches(2.36), "MiniMax 场景图")
    unique_points: List[str] = []
    for product in products[:4]:
        for point in product["selling_points"]:
            if point not in unique_points:
                unique_points.append(point)
    while len(unique_points) < 6:
        unique_points.append("支持包装、颜色、组合方案定制")

    add_card(slide, Inches(0.8), Inches(4.1), Inches(4.2), Inches(2.25), "产品力", unique_points[:4], (37, 99, 235), 9.5)
    add_card(slide, Inches(5.35), Inches(1.48), Inches(3.45), Inches(4.87), "渠道适配", ["商超、电商、礼品、批发渠道", "支持单品陈列和套装组合", "可按目标客户群调整风格", "适合目录册初稿快速沟通"], (14, 165, 233), 9.5)
    add_card(slide, Inches(9.15), Inches(1.48), Inches(3.45), Inches(4.87), "采购沟通", ["确认材质、箱规、MOQ", "补充认证、交期、价格梯度", "导入真实商品图后生成客户版", "可继续按品牌风格重排版"], (16, 185, 129), 9.5)
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def build_presentation(spec: CatalogSpec, pptx_path: Path, run_name: str, assets: List[ImageAsset], now: datetime) -> None:
    products = product_ideas(spec.category)
    assets_by_key = asset_map(assets)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    create_cover(prs, spec, run_name, assets_by_key, now)

    if spec.page_count >= 2:
        create_product_matrix(prs, spec, products, 2, assets_by_key)

    if spec.page_count >= 3:
        create_specs_page(prs, spec, products, 3, assets_by_key)

    page_no = 4
    product_idx = 0
    while page_no <= spec.page_count:
        if page_no == spec.page_count and spec.page_count >= 5:
            create_selling_points(prs, spec, products, page_no, assets_by_key)
        else:
            asset = assets_by_key.get(f"product_{(product_idx % 4) + 1}")
            create_product_detail(prs, spec, products[product_idx % len(products)], page_no, asset)
            product_idx += 1
        page_no += 1

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))


def write_metadata(run_dir: Path, spec: CatalogSpec, result: GenerationResult, now: datetime) -> None:
    metadata = {
        "version": "v1.2",
        "timezone": os.getenv("PPT_TIMEZONE", DEFAULT_TZ),
        "created_at": now.isoformat(timespec="seconds"),
        "spec": asdict(spec),
        "result": asdict(result),
        "output_rule": "PPT 文件直接放在任务目录一级；飞书只回复 windows_path。",
        "image_rule": "生产默认使用 MiniMax image-01；若缺少密钥或生图失败则失败，避免纯文字 PPT。",
    }
    (run_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "README.txt").write_text(
        "OpenClaw PPT Generation v1.2\n"
        f"用户问题：{spec.prompt}\n"
        f"类目：{spec.category}\n"
        f"页数：{spec.page_count}\n"
        f"PPT：{result.windows_path}\n"
        f"图片数量：{result.image_count}\n"
        "飞书回复规则：只回复上面的 Windows/Samba 本地路径。\n",
        encoding="utf-8",
    )


def generate(prompt: str, sender_name: str = "", sender_open_id: str = "", root: Optional[str] = None, image_mode: Optional[str] = None) -> GenerationResult:
    start = time.time()
    if load_dotenv:
        load_dotenv()

    now = now_in_configured_timezone()
    prompt = clean_prompt(prompt)
    default_page_count = env_int("PPT_DEFAULT_PAGE_COUNT", 5)
    max_page_count = env_int("PPT_MAX_PAGE_COUNT", 20)
    page_count = extract_page_count(prompt, default_page_count, max_page_count)
    category = extract_category(prompt)
    spec = CatalogSpec(
        prompt=prompt,
        category=category,
        page_count=page_count,
        doc_type="商品目录册PPT",
        sender_name=sender_name,
        sender_open_id=sender_open_id,
    )

    root_path = Path(root or os.getenv("PPT_CATALOG_ROOT", "/data/share/yaq/ppt/catalog")).expanduser().resolve()
    paths = make_run_paths(prompt, root_path, now)
    run_name = str(paths["run_name"])
    run_dir = Path(paths["run_dir"])
    pptx_path = Path(paths["pptx_path"])
    selected_image_mode = (image_mode or os.getenv("PPT_IMAGE_MODE", "minimax")).strip().lower()

    try:
        products = product_ideas(category)
        image_assets = prepare_images(spec, products, run_dir, selected_image_mode)
        build_presentation(spec, pptx_path, run_name, image_assets, now)
        elapsed = round(time.time() - start, 1)
        windows_path = windows_path_for(pptx_path, root_path)
        result = GenerationResult(
            ok=True,
            prompt=prompt,
            category=category,
            page_count=page_count,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path,
            elapsed_seconds=elapsed,
            image_mode=selected_image_mode,
            image_count=len(image_assets),
            image_assets=[asdict(asset) for asset in image_assets],
            reply_text=windows_path,
        )
        write_metadata(run_dir, spec, result, now)
        return result
    except Exception as exc:
        elapsed = round(time.time() - start, 1)
        windows_path = windows_path_for(pptx_path, root_path)
        return GenerationResult(
            ok=False,
            prompt=prompt,
            category=category,
            page_count=page_count,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path,
            elapsed_seconds=elapsed,
            image_mode=selected_image_mode,
            image_count=0,
            reply_text="",
            error=str(exc),
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate product catalog PPT v1.2")
    parser.add_argument("--prompt", required=True, help="飞书用户原始问题")
    parser.add_argument("--sender-name", default="", help="飞书发送人姓名")
    parser.add_argument("--sender-open-id", default="", help="飞书发送人 open_id")
    parser.add_argument("--root", default=None, help="输出根目录，不填则读 PPT_CATALOG_ROOT")
    parser.add_argument("--image-mode", choices=["minimax", "placeholder", "off"], default=None, help="默认读取 PPT_IMAGE_MODE；生产用 minimax")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出")
    args = parser.parse_args()

    result = generate(args.prompt, args.sender_name, args.sender_open_id, args.root, args.image_mode)
    data = asdict(result)
    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))
    else:
        if result.ok:
            # v1.2 非 JSON 输出只打印本地映射盘路径，便于飞书直接转发。
            print(result.windows_path)
        else:
            print(f"PPT 生成失败：{result.error}", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    main()
