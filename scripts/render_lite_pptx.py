#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Lite fallback PPTX renderer.

This renderer is intentionally conservative. It exists so v1.0 can run even when
upstream PPT-master has not yet been installed. In production, OpenClaw + MiniMax
should generate richer PPT-master/SVG projects using prompt_bundle.md.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any


def _import_pptx():
    try:
        from pptx import Presentation
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
        return Presentation, PP_ALIGN, Inches, Pt
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-pptx 未安装，请先执行 pip install -r requirements.txt") from exc


def _safe_text(value: Any, default: str = "待确认") -> str:
    if value is None:
        return default
    s = str(value).strip()
    return s or default


def _add_title(slide, text: str, Inches, Pt) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(12.0), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.name = "Microsoft YaHei"


def _add_body(slide, lines: list[str], Inches, Pt, x=0.85, y=1.35, w=6.3, h=4.8) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(16)
        p.font.name = "Microsoft YaHei"
        p.space_after = Pt(8)


def _add_image_if_exists(slide, image_path: str | None, Inches, x=7.55, y=1.35, w=4.7, h=4.7) -> None:
    if not image_path:
        return
    path = Path(image_path)
    if not path.exists():
        return
    try:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    except Exception:
        # Do not fail the whole task because of one bad image.
        pass


def _derive_catalog_pages(data: dict) -> list[dict]:
    products = data.get("products") or []
    images = data.get("images") or []
    user_message = _safe_text(data.get("user_message"), "")
    pages: list[dict] = []

    pages.append({"type": "cover", "title": data.get("title") or "商品目录册"})

    if products:
        product = products[0]
        desc = _safe_text(product.get("description"), user_message or "根据用户输入生成商品目录册初稿。")
        pages.append({
            "type": "product_overview",
            "title": "商品概览",
            "lines": [
                f"商品名称：{_safe_text(product.get('name'))}",
                f"商品类别：{_safe_text(product.get('category'))}",
                f"商品描述：{desc}",
            ],
            "image": (product.get("image_paths") or images or [None])[0],
        })

        selling_points = product.get("selling_points") or []
        pages.append({
            "type": "selling_points",
            "title": "建议卖点",
            "lines": [f"• {_safe_text(p, '待根据输入提炼')}" for p in selling_points[:5]] or ["• 待根据输入提炼"],
            "image": (product.get("image_paths") or images or [None])[0],
        })

        specs = product.get("specs") or {}
        pages.append({
            "type": "specs",
            "title": "规格与待确认信息",
            "lines": [
                f"材质：{_safe_text(specs.get('material'))}",
                f"尺寸：{_safe_text(specs.get('size'))}",
                f"颜色：{_safe_text(specs.get('color'))}",
                f"用途：{_safe_text(specs.get('usage'))}",
                f"价格：{_safe_text(specs.get('price'))}",
                f"MOQ：{_safe_text(specs.get('moq'))}",
            ],
            "image": None,
        })

    if len(images) > 1:
        pages.append({"type": "gallery", "title": "图片素材总览", "images": images[:4]})

    pages.append({
        "type": "closing",
        "title": "后续沟通建议",
        "lines": [
            "以上内容为根据飞书输入生成的商品目录册初稿。",
            "规格、价格、材质、MOQ、认证等未明确提供的信息需人工确认。",
            "可继续补充商品参数、品牌信息和目标客户群，以生成更完整版本。",
        ],
    })
    return pages


def render_pptx(product_json_path: Path, output_path: Path) -> Path:
    Presentation, PP_ALIGN, Inches, Pt = _import_pptx()
    data = json.loads(product_json_path.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    pages = _derive_catalog_pages(data)

    for page in pages:
        slide = prs.slides.add_slide(blank)
        title = _safe_text(page.get("title"), "商品目录册")
        _add_title(slide, title, Inches, Pt)
        if page["type"] == "cover":
            _add_body(
                slide,
                [
                    "OpenClaw PPT Generation v1.0",
                    "输入来源：飞书文字/图片",
                    "输出策略：服务器落盘，飞书只回复路径",
                    "说明：商品目录册内容结构可由模型动态规划，本 fallback 仅用于基础可跑通。",
                ],
                Inches,
                Pt,
                x=0.95,
                y=1.65,
                w=8.8,
                h=3.6,
            )
            first_image = (data.get("images") or [None])[0]
            _add_image_if_exists(slide, first_image, Inches, x=9.1, y=1.45, w=3.2, h=3.2)
        elif page["type"] == "gallery":
            imgs = page.get("images") or []
            positions = [(0.9, 1.3), (4.1, 1.3), (7.3, 1.3), (10.5, 1.3)]
            for img, (x, y) in zip(imgs, positions):
                _add_image_if_exists(slide, img, Inches, x=x, y=y, w=2.6, h=2.6)
        else:
            _add_body(slide, page.get("lines") or [], Inches, Pt)
            _add_image_if_exists(slide, page.get("image"), Inches)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    return output_path


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--product-json", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    print(render_pptx(Path(args.product_json), Path(args.output)))
