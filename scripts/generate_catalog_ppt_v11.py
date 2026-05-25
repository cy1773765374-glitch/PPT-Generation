#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
OpenClaw PPT Generation v1.1

功能：
- 从飞书用户问题中解析商品类目和页数
- 生成商品目录册 PPT
- 输出目录：/data/share/yaq/ppt/catalog/YYYY-MM-DD-HHMM-用户询问的问题/
- PPT 文件直接位于任务目录下一级，不再使用 project/exports/
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover
    load_dotenv = None

try:
    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"python-pptx 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)


CN_NUM_MAP = {
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}

INVALID_FILENAME_CHARS = r'<>:"/\\|?*\n\r\t'


@dataclass
class CatalogSpec:
    prompt: str
    category: str
    page_count: int
    doc_type: str
    sender_name: str = ""
    sender_open_id: str = ""


@dataclass
class GenerationResult:
    ok: bool
    prompt: str
    category: str
    page_count: int
    run_name: str
    run_dir: str
    pptx_path: str
    windows_path: str
    elapsed_seconds: float
    error: str = ""


def env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if not value:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def parse_cn_number(text: str) -> Optional[int]:
    text = text.strip()
    if not text:
        return None
    if text.isdigit():
        return int(text)
    if text in CN_NUM_MAP:
        return CN_NUM_MAP[text]
    if text.startswith("十"):
        # 十一、十二
        if len(text) == 1:
            return 10
        return 10 + CN_NUM_MAP.get(text[1:], 0)
    if "十" in text:
        # 二十、二十一
        left, _, right = text.partition("十")
        left_num = CN_NUM_MAP.get(left, 1)
        right_num = CN_NUM_MAP.get(right, 0) if right else 0
        return left_num * 10 + right_num
    return None


def extract_page_count(prompt: str, default: int, max_page_count: int) -> int:
    patterns = [
        r"(\d{1,3})\s*[页頁]",
        r"([一二两三四五六七八九十]{1,3})\s*[页頁]",
        r"(\d{1,3})\s*page",
        r"(\d{1,3})\s*slides?",
    ]
    for pattern in patterns:
        match = re.search(pattern, prompt, flags=re.IGNORECASE)
        if match:
            value = parse_cn_number(match.group(1))
            if value:
                return max(1, min(value, max_page_count))
    return max(1, min(default, max_page_count))


def clean_prompt(prompt: str) -> str:
    prompt = re.sub(r"\s+", " ", prompt.strip())
    # 去除飞书复制时可能带入的回复前缀
    prompt = re.sub(r"^回复\s*[^:：]{1,30}[:：]\s*", "", prompt)
    return prompt.strip()


def extract_category(prompt: str) -> str:
    text = clean_prompt(prompt)

    # 优先抽取 “xxx相关”
    match = re.search(r"(?:生成|做|制作|出|帮我做|帮我生成|设计)?(?:一个|一份)?(.{2,30}?)(?:相关|类|系列)的?", text)
    if match:
        category = match.group(1)
        category = remove_task_words(category)
        if category:
            return category

    # 抽取 “生成一个厨房餐具的3页商品目录册PPT”
    match = re.search(r"(?:生成|做|制作|设计|帮我生成|帮我做)(?:一个|一份)?(.{2,30}?)(?:的)?(?:\d+|[一二两三四五六七八九十]+)?页", text)
    if match:
        category = remove_task_words(match.group(1))
        if category:
            return category

    # 抽取 “xxx商品目录册”
    match = re.search(r"(.{2,30}?)(?:商品目录册|产品目录册|目录册|PPT|ppt)", text)
    if match:
        category = remove_task_words(match.group(1))
        if category:
            return category

    return "综合商品"


def remove_task_words(text: str) -> str:
    words = [
        "生成", "做", "制作", "设计", "帮我", "帮我做", "帮我生成",
        "一个", "一份", "相关", "商品目录册", "产品目录册", "目录册",
        "PPT", "ppt", "幻灯片", "的", "页", "页面",
    ]
    result = text.strip()
    for word in words:
        result = result.replace(word, "")
    result = re.sub(r"\d+", "", result)
    result = re.sub(r"[一二两三四五六七八九十]+", "", result)
    return result.strip(" -_，。,.：:；;")


def sanitize_filename(text: str, max_chars: int = 72) -> str:
    text = clean_prompt(text)
    for ch in INVALID_FILENAME_CHARS:
        text = text.replace(ch, "")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，,。.!！?？：:；;]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-._ ")
    if not text:
        text = "商品目录册"
    if len(text) > max_chars:
        text = text[:max_chars].rstrip("-._ ")
    return text


def make_run_paths(prompt: str, root: Path, now: datetime) -> Dict[str, Path | str]:
    max_chars = env_int("PPT_FILENAME_MAX_CHARS", 72)
    time_prefix = now.strftime("%Y-%m-%d-%H%M")
    slug = sanitize_filename(prompt, max_chars=max_chars)
    base_name = f"{time_prefix}-{slug}"

    auto_dedup = os.getenv("PPT_AUTO_DEDUP", "1") != "0"
    run_name = base_name
    run_dir = root / run_name
    if auto_dedup:
        i = 1
        while run_dir.exists():
            run_name = f"{base_name}-{i:02d}"
            run_dir = root / run_name
            i += 1

    pptx_path = run_dir / f"{run_name}.pptx"
    return {"run_name": run_name, "run_dir": run_dir, "pptx_path": pptx_path}


def windows_path_for(pptx_path: Path, root: Path) -> str:
    prefix = os.getenv("PPT_WINDOWS_SHARE_PREFIX", "Z:\\yaq\\ppt\\catalog").rstrip("\\/")
    try:
        rel = pptx_path.relative_to(root)
    except ValueError:
        rel = pptx_path.name
    rel_win = str(rel).replace("/", "\\")
    return f"{prefix}\\{rel_win}"


def product_ideas(category: str) -> List[Dict[str, Any]]:
    category_key = category.replace(" ", "")
    if any(k in category_key for k in ["厨房餐具", "餐具", "厨具", "厨房"]):
        return [
            {
                "name": "不锈钢刀叉勺套装",
                "position": "家庭日用 / 礼品套装",
                "selling_points": ["食品级不锈钢", "镜面抛光", "多规格组合"],
                "spec": "16件/24件/30件可选；独立彩盒包装",
            },
            {
                "name": "陶瓷碗盘组合",
                "position": "餐桌场景 / 家居零售",
                "selling_points": ["高温烧制", "釉面细腻", "适配多风格餐桌"],
                "spec": "碗、盘、汤碗、调味碟组合；颜色可定制",
            },
            {
                "name": "硅胶厨具铲勺套装",
                "position": "厨房烹饪 / 防刮锅具",
                "selling_points": ["耐高温硅胶", "不伤锅涂层", "易清洗收纳"],
                "spec": "5件/8件/12件套；可加收纳桶",
            },
            {
                "name": "厨房餐具收纳架",
                "position": "厨房收纳 / 台面整理",
                "selling_points": ["分区收纳", "沥水设计", "节省台面空间"],
                "spec": "单层/双层；金属/塑料材质可选",
            },
        ]
    return [
        {
            "name": f"{category}核心款 A",
            "position": "主推款 / 常规渠道",
            "selling_points": ["高频使用", "稳定供应", "支持定制"],
            "spec": "尺寸、颜色、包装、MOQ 建议确认",
        },
        {
            "name": f"{category}升级款 B",
            "position": "中高端款 / 礼品渠道",
            "selling_points": ["质感升级", "外观差异化", "适合套装销售"],
            "spec": "材质、工艺、认证、价格建议确认",
        },
        {
            "name": f"{category}组合款 C",
            "position": "套装款 / 批量采购",
            "selling_points": ["组合灵活", "性价比高", "便于陈列"],
            "spec": "组合数量、箱规、交期建议确认",
        },
    ]


def set_slide_background(slide, rgb=(248, 250, 252)) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def add_textbox(slide, left, top, width, height, text, font_size=20, bold=False, color=(31, 41, 55), align=None):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = "Microsoft YaHei"
    return box


def add_title(slide, title: str, subtitle: str = "") -> None:
    add_textbox(slide, Inches(0.65), Inches(0.38), Inches(12.0), Inches(0.55), title, 28, True, (15, 23, 42))
    if subtitle:
        add_textbox(slide, Inches(0.68), Inches(0.92), Inches(11.8), Inches(0.35), subtitle, 11, False, (100, 116, 139))
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(1.28), Inches(12.0), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(37, 99, 235)
    line.line.fill.background()


def add_card(slide, left, top, width, height, title: str, body_lines: List[str], accent=(37, 99, 235)) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    shape.line.color.rgb = RGBColor(226, 232, 240)
    shape.line.width = Pt(1)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()

    add_textbox(slide, left + Inches(0.22), top + Inches(0.16), width - Inches(0.35), Inches(0.32), title, 15, True, (15, 23, 42))
    body = "\n".join([f"• {line}" for line in body_lines])
    box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.55), width - Inches(0.38), height - Inches(0.65))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(10.5)
        p.font.name = "Microsoft YaHei"
        p.font.color.rgb = RGBColor(71, 85, 105)
        p.space_after = Pt(3)


def add_footer(slide, text: str) -> None:
    add_textbox(slide, Inches(0.68), Inches(7.05), Inches(12.0), Inches(0.25), text, 8.5, False, (148, 163, 184))


def create_cover(prs: Presentation, spec: CatalogSpec, run_name: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, (241, 245, 249))

    # 顶部装饰块
    hero = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(0.65), Inches(12.0), Inches(5.7))
    hero.fill.solid()
    hero.fill.fore_color.rgb = RGBColor(255, 255, 255)
    hero.line.color.rgb = RGBColor(226, 232, 240)

    add_textbox(slide, Inches(1.15), Inches(1.25), Inches(11.0), Inches(0.65), f"{spec.category}商品目录册", 34, True, (15, 23, 42))
    add_textbox(slide, Inches(1.18), Inches(2.02), Inches(10.6), Inches(0.45), "Product Catalog · Feishu Generated Draft", 16, False, (37, 99, 235))
    add_textbox(slide, Inches(1.18), Inches(2.78), Inches(10.8), Inches(0.82), f"输入需求：{spec.prompt}", 15, False, (51, 65, 85))

    meta = [
        f"页数：{spec.page_count} 页",
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"生成版本：OpenClaw PPT Generation v1.1",
    ]
    if spec.sender_name:
        meta.append(f"发起人：{spec.sender_name}")
    add_card(slide, Inches(1.18), Inches(4.08), Inches(5.45), Inches(1.45), "目录册信息", meta, (37, 99, 235))
    add_card(slide, Inches(6.92), Inches(4.08), Inches(4.95), Inches(1.45), "交付说明", ["PPT 文件直接位于任务目录一级", "可通过 Samba 映射盘打开", "内容为第一版可编辑草稿"], (14, 165, 233))

    add_footer(slide, run_name)


def create_product_matrix(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "商品系列矩阵", f"围绕“{spec.category}”生成的第一版目录册结构")

    positions = [
        (Inches(0.75), Inches(1.62)),
        (Inches(6.85), Inches(1.62)),
        (Inches(0.75), Inches(4.1)),
        (Inches(6.85), Inches(4.1)),
    ]
    for idx, product in enumerate(products[:4]):
        left, top = positions[idx]
        lines = [
            f"定位：{product['position']}",
            f"卖点：{' / '.join(product['selling_points'])}",
            f"规格：{product['spec']}",
        ]
        add_card(slide, left, top, Inches(5.75), Inches(2.05), product["name"], lines, (37, 99, 235) if idx % 2 == 0 else (14, 165, 233))
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_selling_points(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "核心卖点与采购价值", "用于销售沟通、客户初筛、目录册初版展示")

    points = []
    for product in products[:3]:
        points.extend(product["selling_points"][:2])
    unique_points = []
    for p in points:
        if p not in unique_points:
            unique_points.append(p)
    while len(unique_points) < 6:
        unique_points.append("支持包装、颜色、组合方案定制")

    lefts = [Inches(0.8), Inches(4.75), Inches(8.7)]
    titles = ["产品力", "渠道适配", "采购沟通"]
    bodies = [
        unique_points[:3],
        ["适合商超、电商、礼品与批发渠道", "支持单品陈列和套装组合", "可按目标客户群调整风格"],
        ["建议确认材质、箱规、MOQ", "建议补充认证、交期、价格梯度", "可继续输入品牌信息生成正式版"],
    ]
    for i in range(3):
        add_card(slide, lefts[i], Inches(1.75), Inches(3.55), Inches(4.7), titles[i], bodies[i], [(37,99,235),(14,165,233),(16,185,129)][i])
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_specs_page(prs: Presentation, spec: CatalogSpec, products: List[Dict[str, Any]], page_no: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, "规格建议与待确认信息", "便于业务人员继续补充真实商品参数")

    headers = ["项目", "建议内容", "备注"]
    rows = [
        ["商品系列", " / ".join([p["name"] for p in products[:3]]), "可根据实际 SKU 调整"],
        ["材质", "不锈钢 / 陶瓷 / 食品级硅胶 / PP 等", "按具体产品确认"],
        ["包装", "彩盒、礼盒、裸装、外箱", "建议补充包装尺寸"],
        ["定制", "LOGO、颜色、组合数量、说明书", "确认 MOQ 与打样周期"],
        ["贸易信息", "价格、箱规、交期、认证", "需业务人工补充"],
    ]

    x = Inches(0.8)
    y = Inches(1.65)
    table = slide.shapes.add_table(len(rows) + 1, 3, x, y, Inches(11.8), Inches(4.4)).table
    table.columns[0].width = Inches(1.55)
    table.columns[1].width = Inches(6.8)
    table.columns[2].width = Inches(3.45)

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(37, 99, 235)
        for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.name = "Microsoft YaHei"

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(255, 255, 255)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(10)
                p.font.color.rgb = RGBColor(51, 65, 85)
                p.font.name = "Microsoft YaHei"

    add_card(slide, Inches(0.8), Inches(6.22), Inches(11.8), Inches(0.72), "下一步建议", ["补充真实图片、价格、材质、MOQ、认证后，可生成正式客户版目录册。"], (16, 185, 129))
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def create_product_detail(prs: Presentation, spec: CatalogSpec, product: Dict[str, Any], page_no: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_title(slide, product["name"], f"{spec.category} · 单品详情页")

    # 左侧模拟图片区，占位但可编辑
    img_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.85), Inches(1.62), Inches(4.5), Inches(4.8))
    img_box.fill.solid()
    img_box.fill.fore_color.rgb = RGBColor(226, 232, 240)
    img_box.line.color.rgb = RGBColor(203, 213, 225)
    add_textbox(slide, Inches(1.25), Inches(3.65), Inches(3.7), Inches(0.5), "商品图片占位\n可替换为实物图/AI图", 16, True, (100, 116, 139), PP_ALIGN.CENTER)

    add_card(slide, Inches(5.75), Inches(1.62), Inches(6.75), Inches(1.55), "产品定位", [product["position"], "适合用于目录册主推款展示"], (37, 99, 235))
    add_card(slide, Inches(5.75), Inches(3.42), Inches(6.75), Inches(1.55), "核心卖点", product["selling_points"], (14, 165, 233))
    add_card(slide, Inches(5.75), Inches(5.22), Inches(6.75), Inches(1.15), "规格建议", [product["spec"]], (16, 185, 129))
    add_footer(slide, f"{spec.category}商品目录册 · 第 {page_no} 页")


def build_presentation(spec: CatalogSpec, pptx_path: Path, run_name: str) -> None:
    products = product_ideas(spec.category)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 删除默认空白演示文稿中的所有 slide：python-pptx 新建时默认没有 slide，这里保留安全写法。
    create_cover(prs, spec, run_name)

    if spec.page_count >= 2:
        create_product_matrix(prs, spec, products, 2)

    if spec.page_count >= 3:
        create_specs_page(prs, spec, products, 3)

    # 4页及以上，增加单品详情/卖点页
    page_no = 4
    product_idx = 0
    while page_no <= spec.page_count:
        if page_no == spec.page_count and spec.page_count >= 5:
            create_selling_points(prs, spec, products, page_no)
        else:
            create_product_detail(prs, spec, products[product_idx % len(products)], page_no)
            product_idx += 1
        page_no += 1

    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))


def write_metadata(run_dir: Path, spec: CatalogSpec, result: GenerationResult) -> None:
    metadata = {
        "version": "v1.1",
        "spec": asdict(spec),
        "result": asdict(result),
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "output_rule": "PPT 文件直接放在任务目录一级，不使用 project/exports。",
    }
    (run_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "README.txt").write_text(
        "OpenClaw PPT Generation v1.1\n"
        f"用户问题：{spec.prompt}\n"
        f"类目：{spec.category}\n"
        f"页数：{spec.page_count}\n"
        f"PPT：{result.pptx_path}\n"
        f"Windows 路径：{result.windows_path}\n",
        encoding="utf-8",
    )


def generate(prompt: str, sender_name: str = "", sender_open_id: str = "", root: Optional[str] = None) -> GenerationResult:
    start = time.time()
    if load_dotenv:
        load_dotenv()

    prompt = clean_prompt(prompt)
    default_page_count = env_int("PPT_DEFAULT_PAGE_COUNT", 5)
    max_page_count = env_int("PPT_MAX_PAGE_COUNT", 20)
    page_count = extract_page_count(prompt, default_page_count, max_page_count)
    category = extract_category(prompt)
    spec = CatalogSpec(
        prompt=prompt,
        category=category,
        page_count=page_count,
        doc_type="商品目录册PPT",
        sender_name=sender_name,
        sender_open_id=sender_open_id,
    )

    root_path = Path(root or os.getenv("PPT_CATALOG_ROOT", "/data/share/yaq/ppt/catalog")).expanduser().resolve()
    paths = make_run_paths(prompt, root_path, datetime.now())
    run_name = str(paths["run_name"])
    run_dir = Path(paths["run_dir"])
    pptx_path = Path(paths["pptx_path"])

    try:
        build_presentation(spec, pptx_path, run_name)
        elapsed = round(time.time() - start, 1)
        result = GenerationResult(
            ok=True,
            prompt=prompt,
            category=category,
            page_count=page_count,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path_for(pptx_path, root_path),
            elapsed_seconds=elapsed,
        )
        write_metadata(run_dir, spec, result)
        return result
    except Exception as exc:
        elapsed = round(time.time() - start, 1)
        return GenerationResult(
            ok=False,
            prompt=prompt,
            category=category,
            page_count=page_count,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path_for(pptx_path, root_path),
            elapsed_seconds=elapsed,
            error=str(exc),
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate product catalog PPT v1.1")
    parser.add_argument("--prompt", required=True, help="飞书用户原始问题")
    parser.add_argument("--sender-name", default="", help="飞书发送人姓名")
    parser.add_argument("--sender-open-id", default="", help="飞书发送人 open_id")
    parser.add_argument("--root", default=None, help="输出根目录，不填则读 PPT_CATALOG_ROOT")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出")
    args = parser.parse_args()

    result = generate(args.prompt, args.sender_name, args.sender_open_id, args.root)
    data = asdict(result)
    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))
    else:
        if result.ok:
            print(f"已完成，PPT 文件：\n{result.windows_path}\n")
            print(f"服务器路径：\n{result.pptx_path}\n")
            print(f"已完成 · {result.page_count}页 · 耗时 {result.elapsed_seconds}s")
        else:
            print(f"PPT 生成失败：{result.error}", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    main()
