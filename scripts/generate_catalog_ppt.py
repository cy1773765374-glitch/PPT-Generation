#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
OpenClaw PPT 商品目录册生成器（PPT-master 适配版）

设计目标：
- 用户显式页数、语言、品牌、页面结构、产品类目优先。
- deck_plan.json 是唯一事实来源，渲染层不得重新决定页数或类目。
- 默认优先使用 vendor/ppt-master 的 SVG → native editable PPTX 管线。
- 生产默认调用 MiniMax 生图，失败即失败；离线测试可用 placeholder。
- 成功时飞书只回复 Windows/Samba 本地路径。
"""

from __future__ import annotations

import argparse
import base64
import json
import os
import re
import socket
import sys
import time
import urllib.parse
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from zoneinfo import ZoneInfo

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover
    load_dotenv = None

try:
    import requests
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"requests 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)

try:
    from PIL import Image, ImageDraw, ImageFont
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"Pillow 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except Exception as exc:  # pragma: no cover
    print(json.dumps({"ok": False, "error": f"python-pptx 未安装或加载失败：{exc}"}, ensure_ascii=False))
    sys.exit(1)

try:
    from openclaw_pptmaster_adapter import (
        PPTMasterPipelineError,
        PPTMasterUnavailable,
        build_with_pptmaster,
    )
except Exception:  # pragma: no cover
    PPTMasterPipelineError = RuntimeError  # type: ignore
    PPTMasterUnavailable = RuntimeError  # type: ignore
    build_with_pptmaster = None  # type: ignore


DEFAULT_TZ = "Asia/Shanghai"
INVALID_FILENAME_CHARS = '<>:"/\\\\|?*\n\r\t'
CJK_RE = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff\uf900-\ufaff]")

CN_NUM_MAP = {
    "零": 0,
    "一": 1,
    "二": 2,
    "两": 2,
    "三": 3,
    "四": 4,
    "五": 5,
    "六": 6,
    "七": 7,
    "八": 8,
    "九": 9,
    "十": 10,
}


@dataclass
class ImageBrief:
    key: str
    page: int
    title: str
    prompt: str
    aspect_ratio: str = "16:9"


@dataclass
class SlidePlan:
    page: int
    layout: str
    title: str
    subtitle: str = ""
    category: str = ""
    bullets: List[str] = field(default_factory=list)
    sections: List[Dict[str, Any]] = field(default_factory=list)
    image_key: str = ""
    image_prompt: str = ""
    layout_variant: str = ""


@dataclass
class DeckPlan:
    prompt: str
    language: str
    page_count: int
    deck_type: str
    brand_or_company: str
    theme: str
    style: str
    mode: str
    slides: List[SlidePlan]
    explicit_categories: List[str] = field(default_factory=list)
    image_briefs: List[ImageBrief] = field(default_factory=list)
    style_family: str = "commercial"
    requested_sections: List[str] = field(default_factory=list)


@dataclass
class ImageAsset:
    key: str
    title: str
    prompt: str
    path: str
    source: str
    aspect_ratio: str
    page: int


@dataclass
class ValidationResult:
    ok: bool
    checks: List[str] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)


@dataclass
class GenerationResult:
    ok: bool
    prompt: str
    page_count: int
    run_name: str
    run_dir: str
    pptx_path: str
    windows_path: str
    elapsed_seconds: float
    language: str = ""
    deck_plan_path: str = ""
    image_mode: str = ""
    render_backend: str = ""
    pptmaster_project_dir: str = ""
    pptmaster_log_path: str = ""
    image_count: int = 0
    image_assets: List[Dict[str, Any]] = field(default_factory=list)
    validation: Dict[str, Any] = field(default_factory=dict)
    reply_text: str = ""
    error: str = ""


def env_int(name: str, default: int) -> int:
    value = os.getenv(name)
    if not value:
        return default
    try:
        return int(value)
    except ValueError:
        return default


def env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or value == "":
        return default
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


def now_in_configured_timezone() -> datetime:
    tz_name = os.getenv("PPT_TIMEZONE", DEFAULT_TZ).strip() or DEFAULT_TZ
    try:
        return datetime.now(ZoneInfo(tz_name))
    except Exception:
        return datetime.now(ZoneInfo(DEFAULT_TZ))


def parse_cn_number(text: str) -> Optional[int]:
    text = text.strip()
    if not text:
        return None
    if text.isdigit():
        return int(text)
    if text in CN_NUM_MAP:
        return CN_NUM_MAP[text]
    if text.startswith("十"):
        if len(text) == 1:
            return 10
        return 10 + CN_NUM_MAP.get(text[1:], 0)
    if "十" in text:
        left, _, right = text.partition("十")
        left_num = CN_NUM_MAP.get(left, 1)
        right_num = CN_NUM_MAP.get(right, 0) if right else 0
        return left_num * 10 + right_num
    return None


def clean_prompt(prompt: str) -> str:
    prompt = re.sub(r"\s+", " ", (prompt or "").strip())
    prompt = re.sub(r"^回复\s*[^:：]{1,30}[:：]\s*", "", prompt)
    for marker in ["已完成，PPT 文件", "服务器路径", "已完成 ·", "PPT 文件："]:
        if marker in prompt:
            prompt = prompt.split(marker, 1)[0].strip()
    return prompt.strip()


def extract_page_count(prompt: str, default: int = 5, max_page_count: int = 20) -> int:
    text = clean_prompt(prompt)
    patterns = [
        r"(\d{1,3})\s*[-–—]?\s*page\b",
        r"(\d{1,3})\s*pages\b",
        r"(\d{1,3})\s*slides?\b",
        r"(\d{1,3})\s*[页頁]",
        r"([一二两三四五六七八九十]{1,3})\s*[页頁]",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            value = parse_cn_number(match.group(1))
            if value:
                return max(1, min(value, max_page_count))

    # 明确 Page 1 / Pages 2-10 但开头没写 10-page 时，用最大页码兜底。
    max_page = 0
    for m in re.finditer(r"Pages?\s+(\d{1,3})(?:\s*[-–—]\s*(\d{1,3}))?", text, flags=re.IGNORECASE):
        right = m.group(2) or m.group(1)
        max_page = max(max_page, int(right))
    if max_page:
        return max(1, min(max_page, max_page_count))
    return max(1, min(default, max_page_count))


def detect_language(prompt: str) -> str:
    text = clean_prompt(prompt)
    lower = text.lower()
    if re.search(r"all\s+text\s+in\s+english|english\s+descriptions?|in\s+english|英文|英语", lower, flags=re.IGNORECASE):
        return "en"
    if CJK_RE.search(text):
        return "zh"
    return "en"


def extract_brand_or_company(prompt: str, language: str) -> str:
    text = clean_prompt(prompt)
    patterns = [
        r"for\s+([A-Z][A-Z0-9& .'-]{2,60}?)(?:[\.,;]|\s+Page\s+\d|\s+Pages\s+\d|$)",
        r"Company\s+introduction\s+for\s+([A-Z][A-Z0-9& .'-]{2,60}?)(?:\s*[-–—:]|[\.,;]|$)",
        r"brand\s*[:：]\s*([^\.;，。]{2,60})",
        r"company\s*[:：]\s*([^\.;，。]{2,60})",
        r"公司\s*[:：]\s*([^\.;，。]{2,60})",
        r"品牌\s*[:：]\s*([^\.;，。]{2,60})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            value = re.sub(r"\s+", " ", match.group(1)).strip(" -–—:：,，.;。")
            # 避免把 Product Catalog PPT 误吸进品牌名。
            value = re.sub(r"\s+PPT$", "", value, flags=re.IGNORECASE).strip()
            if value:
                return value
    return "SELLERS UNION" if language == "en" and "SELLERS UNION" in text.upper() else ("Your Company" if language == "en" else "品牌公司")


def extract_style(prompt: str, language: str) -> str:
    text = clean_prompt(prompt)
    patterns = [
        r"(High-quality[^\.。;；]{2,120}style)",
        r"(premium[^\.。;；]{2,120}style)",
        r"风格\s*[:：]\s*([^\.。;；]{2,120})",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip(" .。;；")
    return "High-quality commercial product catalog style" if language == "en" else "高质量商业商品目录册风格"


def extract_theme(prompt: str, language: str) -> str:
    text = clean_prompt(prompt)
    match = re.search(r"Generate\s+a\s+\d{1,3}\s*[-–—]?\s*page\s+(.+?)\s+product\s+catalog", text, flags=re.IGNORECASE)
    if match:
        return re.sub(r"\s+", " ", match.group(1)).strip(" .;:")
    match = re.search(r"(.{2,40}?)(?:相关|主题|类|系列)?(?:的)?(?:\d+|[一二两三四五六七八九十]+)?页?(?:商品目录册|产品目录册|目录册|PPT|ppt)", text)
    if match:
        raw = match.group(1)
        raw = re.sub(r"^(生成|做|制作|设计|帮我|帮我做|帮我生成|一个|一份)+", "", raw).strip()
        raw = re.sub(r"\d+|[一二两三四五六七八九十]+", "", raw).strip(" -_，。,.：:；;")
        if raw:
            return raw
    return "product catalog" if language == "en" else "商品目录册"


def title_case_category(text: str) -> str:
    text = re.sub(r"\s+", " ", text.strip(" -–—:：,，.;。"))
    if not text:
        return text
    # 用户已经给了全大写缩写时保留，例如 COS。
    words = []
    for w in text.split(" "):
        if w.isupper() and len(w) <= 5:
            words.append(w)
        elif w.lower() in {"and", "or", "of", "for", "with"}:
            words.append(w.lower())
        else:
            words.append(w[:1].upper() + w[1:])
    if words:
        words[0] = words[0][:1].upper() + words[0][1:]
    return " ".join(words)




STYLE_PROFILES: Dict[str, Dict[str, Any]] = {
    "commercial": {
        "name_en": "Modern Commercial",
        "name_zh": "现代商业",
        "bg": (248, 250, 252),
        "panel": (255, 255, 255),
        "accent": (37, 99, 235),
        "accent2": (14, 165, 233),
        "accent3": (16, 185, 129),
        "text": (15, 23, 42),
        "muted": (71, 85, 105),
    },
    "minimal": {
        "name_en": "Minimal Editorial",
        "name_zh": "极简画册",
        "bg": (250, 250, 249),
        "panel": (255, 255, 255),
        "accent": (30, 41, 59),
        "accent2": (100, 116, 139),
        "accent3": (148, 163, 184),
        "text": (17, 24, 39),
        "muted": (82, 82, 91),
    },
    "luxury": {
        "name_en": "Premium Luxury",
        "name_zh": "高端轻奢",
        "bg": (250, 247, 242),
        "panel": (255, 252, 247),
        "accent": (146, 64, 14),
        "accent2": (180, 83, 9),
        "accent3": (217, 119, 6),
        "text": (28, 25, 23),
        "muted": (87, 83, 78),
    },
    "playful": {
        "name_en": "Playful Retail",
        "name_zh": "活泼零售",
        "bg": (255, 247, 237),
        "panel": (255, 255, 255),
        "accent": (234, 88, 12),
        "accent2": (217, 70, 239),
        "accent3": (14, 165, 233),
        "text": (30, 41, 59),
        "muted": (71, 85, 105),
    },
    "festive": {
        "name_en": "Festive Campaign",
        "name_zh": "节庆活动",
        "bg": (255, 251, 235),
        "panel": (255, 255, 255),
        "accent": (220, 38, 38),
        "accent2": (245, 158, 11),
        "accent3": (22, 163, 74),
        "text": (39, 39, 42),
        "muted": (82, 82, 91),
    },
    "nature": {
        "name_en": "Natural Warmth",
        "name_zh": "自然温润",
        "bg": (246, 248, 240),
        "panel": (255, 255, 251),
        "accent": (76, 120, 86),
        "accent2": (132, 111, 74),
        "accent3": (101, 163, 13),
        "text": (31, 41, 55),
        "muted": (75, 85, 99),
    },
    "tech": {
        "name_en": "Clean Tech",
        "name_zh": "科技理性",
        "bg": (239, 246, 255),
        "panel": (255, 255, 255),
        "accent": (29, 78, 216),
        "accent2": (6, 182, 212),
        "accent3": (99, 102, 241),
        "text": (15, 23, 42),
        "muted": (51, 65, 85),
    },
}

SECTION_LABELS = {
    "en": {
        "positioning": "Product Positioning",
        "range": "Range Direction",
        "scenario": "Usage Scenario",
        "selling_points": "Selling Points",
        "materials": "Material & Finish",
        "colors": "Color Direction",
        "customization": "Customization",
        "packaging": "Packaging",
        "price": "Price Tier",
        "moq": "MOQ Note",
        "carton": "Carton / Packing Data",
        "certification": "Certification",
        "delivery": "Delivery Note",
        "sku": "SKU Ideas",
        "audience": "Buyer Angle",
    },
    "zh": {
        "positioning": "产品定位",
        "range": "系列方向",
        "scenario": "使用场景",
        "selling_points": "核心卖点",
        "materials": "材质与工艺",
        "colors": "色彩方向",
        "customization": "定制方向",
        "packaging": "包装方式",
        "price": "价格层级",
        "moq": "MOQ 备注",
        "carton": "箱规/装箱信息",
        "certification": "认证信息",
        "delivery": "交期说明",
        "sku": "SKU 建议",
        "audience": "采购视角",
    },
}

REQUESTED_SECTION_PATTERNS: Dict[str, List[str]] = {
    "scenario": ["scenario", "occasion", "use case", "应用场景", "适用场景", "使用场景", "场景"],
    "selling_points": ["selling point", "selling points", "key feature", "key features", "product feature", "product features", "advantage", "卖点", "特点", "优势", "亮点"],
    "materials": ["material", "finish", "材质", "材料", "工艺", "表面处理"],
    "colors": ["color", "colour", "palette", "颜色", "色彩", "配色"],
    "customization": ["custom", "oem", "odm", "定制", "贴牌", "定做"],
    "packaging": ["packaging", "packing", "package", "包装"],
    "price": ["price", "pricing", "报价", "价格", "价位"],
    "moq": ["moq", "minimum order", "起订", "起订量"],
    "carton": ["carton", "ctn", "箱规", "装箱", "外箱"],
    "certification": ["certification", "certificate", "certified", "认证", "证书", "检测"],
    "delivery": ["delivery", "lead time", "shipping", "交期", "货期", "发货"],
    "sku": ["sku", "item no", "货号", "款号", "单品", "产品编号"],
}

CATEGORY_PRESETS_EN: Dict[str, Dict[str, Any]] = {
    "notebook": {
        "positioning": ["Everyday writing and study notebooks for school, office and gifting programs."],
        "range": ["Spiral notebooks", "Soft-cover journals", "Subject notebooks", "Pocket memo books"],
        "scenario": ["Back-to-school sets, office stationery shelves, promotional bundles"],
        "selling_points": ["Clean cover direction", "Flexible paper ruling", "Easy to build coordinated stationery sets"],
        "materials": ["Paper cover, PP cover, kraft cover, inner pages by requested GSM"],
    },
    "pen": {
        "positioning": ["High-frequency writing items suited to retail multipacks and school supply programs."],
        "range": ["Gel pens", "Ballpoint pens", "Color pens", "Mechanical pencils and refill sets"],
        "scenario": ["Daily writing, exam stationery, office desks, gift stationery packs"],
        "selling_points": ["Smooth writing feel", "Color and tip-size options", "Good bundle compatibility"],
        "materials": ["Plastic barrel, soft grip, metal clip options, refillable structures when requested"],
    },
    "eraser": {
        "positioning": ["Compact correction essentials for student stationery and value packs."],
        "range": ["PVC-free erasers", "Novelty shaped erasers", "Dust-free erasers", "Eraser multipacks"],
        "scenario": ["School lists, exam kits, checkout add-ons, children’s stationery ranges"],
        "selling_points": ["Soft erasing feel", "Low residue direction", "Colorful shapes for retail appeal"],
    },
    "ruler": {
        "positioning": ["Measuring tools for school kits, office drawers and drawing sets."],
        "range": ["Plastic rulers", "Flexible rulers", "Geometry sets", "Transparent measuring tools"],
        "scenario": ["Classroom use, exam preparation, drafting, student value packs"],
        "selling_points": ["Clear scale visibility", "Lightweight structure", "Easy set combination"],
    },
    "correction": {
        "positioning": ["Practical correction tools for school, office and study desks."],
        "range": ["Mini correction tape", "Ergonomic tape", "Refillable tape", "Multi-pack correction sets"],
        "scenario": ["Homework correction, office paperwork, exam preparation, desk stationery"],
        "selling_points": ["Clean coverage", "Portable size", "Good for multi-pack retail"],
    },
    "school bag": {
        "positioning": ["Daily carry products designed for school, travel and student lifestyle assortments."],
        "range": ["Backpacks", "Lunch bags", "Drawstring bags", "Lightweight student bags"],
        "scenario": ["Back-to-school programs, student travel, campus retail, gift bundles"],
        "selling_points": ["Comfortable carrying direction", "Compartment planning", "Themeable exterior design"],
        "materials": ["Polyester, nylon, oxford fabric, padded straps when required"],
    },
    "pen case": {
        "positioning": ["Storage accessories for coordinated stationery collections and school sets."],
        "range": ["Zipper pencil cases", "EVA cases", "Transparent pouches", "Multi-compartment cases"],
        "scenario": ["Student desks, school bags, gift stationery sets, retail shelf programs"],
        "selling_points": ["Compact storage", "Easy color matching", "Good add-on category for stationery ranges"],
    },
    "paint": {
        "positioning": ["Creative art supplies for school projects, hobby painting and children’s activity programs."],
        "range": ["Watercolor sets", "Acrylic paint sets", "Poster paints", "Brush and palette kits"],
        "scenario": ["Art classes, craft activities, creative gifts, seasonal DIY programs"],
        "selling_points": ["Bright color presentation", "Set-based selling", "Easy pairing with brushes and paper"],
        "materials": ["Water-based formulas and packaging details can be specified by real product data"],
    },
    "sharpener": {
        "positioning": ["Small desk essentials for school lists, pencil kits and checkout displays."],
        "range": ["Single-hole sharpeners", "Double-hole sharpeners", "Container sharpeners", "Novelty sharpeners"],
        "scenario": ["Classroom use, pencil kits, retail counters, children’s stationery sets"],
        "selling_points": ["Portable size", "Simple replacement SKU", "Works well in multi-piece stationery bundles"],
    },
    "balloon": {
        "positioning": ["Core visual decoration items for party programs and seasonal celebration ranges."],
        "range": ["Latex balloons", "Foil balloons", "Number balloons", "Balloon arch kits"],
        "scenario": ["Birthday parties, weddings, retail party kits, seasonal events"],
        "selling_points": ["Coordinated color sets", "Strong visual impact", "Easy bundle planning"],
    },
    "tableware": {
        "positioning": ["Disposable table setup products for party, picnic and event catering assortments."],
        "range": ["Paper plates", "Cups", "Napkins", "Straws", "Cutlery sets", "Table covers"],
        "scenario": ["Party tables, buffet setup, outdoor events, themed retail kits"],
        "selling_points": ["Complete table solution", "Color-matched merchandising", "Easy cleanup for end users"],
    },
}

CATEGORY_PRESETS_ZH: Dict[str, Dict[str, Any]] = {
    "笔记本": {
        "positioning": ["面向学习、办公和礼品组合的高频书写类产品。"],
        "range": ["线圈本", "软抄本", "主题本", "便携记事本"],
        "scenario": ["开学季、办公采购、文具礼盒、零售陈列"],
        "selling_points": ["封面风格可延展", "内页规格灵活", "适合成套销售"],
    },
    "笔": {
        "positioning": ["适合学校、办公和组合包的基础书写工具。"],
        "range": ["中性笔", "圆珠笔", "彩色笔", "自动铅笔及替芯"],
        "scenario": ["日常书写、考试文具、办公桌面、礼品文具包"],
        "selling_points": ["书写顺滑", "颜色和笔尖规格可选", "便于做多支装"],
    },
    "橡皮": {
        "positioning": ["学生文具和组合套装中的基础修正类单品。"],
        "range": ["无 PVC 橡皮", "造型橡皮", "少屑橡皮", "多枚装橡皮"],
        "scenario": ["学生清单、考试套装、收银台加购、儿童文具系列"],
        "selling_points": ["擦除体验柔和", "造型和颜色适合零售展示", "适合小包装组合"],
    },
    "尺": {
        "positioning": ["适合学生、办公和绘图套装的测量工具。"],
        "range": ["直尺", "软尺", "几何套尺", "透明测量工具"],
        "scenario": ["课堂、考试准备、绘图、学生套装"],
        "selling_points": ["刻度清晰", "结构轻便", "容易和文具套装组合"],
    },
    "书包": {
        "positioning": ["面向学生日常通勤、校园生活和开学季项目的背包类产品。"],
        "range": ["双肩包", "午餐包", "抽绳包", "轻便学生包"],
        "scenario": ["开学季、学生出行、校园零售、礼品组合"],
        "selling_points": ["背负舒适", "分层收纳", "外观主题可延展"],
    },
}


def detect_style_family(prompt: str, theme: str = "", language: str = "en") -> str:
    text = f"{prompt} {theme}".lower()
    checks = [
        ("minimal", ["minimal", "clean", "simple", "editorial", "极简", "简洁", "留白", "画册"]),
        ("luxury", ["luxury", "premium", "high-end", "elegant", "高级", "高端", "轻奢", "质感"]),
        ("playful", ["cute", "kids", "playful", "colorful", "cartoon", "children", "可爱", "儿童", "活泼", "卡通", "彩色"]),
        ("festive", ["festive", "party", "holiday", "christmas", "celebration", "节庆", "派对", "圣诞", "庆典", "节日"]),
        ("nature", ["natural", "eco", "organic", "warm", "wood", "green", "自然", "环保", "温润", "木质", "绿色"]),
        ("tech", ["tech", "digital", "smart", "industrial", "科技", "智能", "数码", "工业", "蓝色科技"]),
    ]
    for family, keys in checks:
        if any(k in text for k in keys):
            return family
    return "commercial"


def style_profile_for(plan: DeckPlan | str) -> Dict[str, Any]:
    family = plan if isinstance(plan, str) else getattr(plan, "style_family", "commercial")
    return STYLE_PROFILES.get(family, STYLE_PROFILES["commercial"])


def extract_page1_description(prompt: str) -> str:
    text = clean_prompt(prompt)
    match = re.search(r"Page\s*1\s*[:：]\s*(.+?)(?:\s+Pages?\s+\d|\s+Page\s+2\s*[:：]|$)", text, flags=re.IGNORECASE)
    if match:
        return match.group(1).strip(" .。;；")
    match = re.search(r"(?:公司介绍|企业介绍)\s*[:：-]?\s*([^。.;；]{4,120})", text, flags=re.IGNORECASE)
    if match:
        return match.group(1).strip(" .。;；")
    return ""


def strip_negated_requirements(text: str) -> str:
    # 先移除“不要/不需要/no/without ...”这一类否定字段，避免用户说“不要价格 MOQ 箱规”时反而被识别成必填字段。
    text = re.sub(r"(?:不要|不需要|无需|不用|不含|不写|别写|去掉)[^。.;；,，]{0,60}", " ", text, flags=re.IGNORECASE)
    text = re.sub(r"\b(?:no|not|without|exclude|remove)\b[^。.;；,，]{0,60}", " ", text, flags=re.IGNORECASE)
    return text


def extract_requested_sections(prompt: str) -> List[str]:
    raw = clean_prompt(prompt)
    lower_raw = raw.lower()
    if any(k in lower_raw for k in ["只要产品方向", "只放产品方向", "only product positioning", "product direction only"]):
        return ["positioning"]
    text = strip_negated_requirements(lower_raw)
    requested: List[str] = []
    for key, patterns in REQUESTED_SECTION_PATTERNS.items():
        if any(p.lower() in text for p in patterns):
            requested.append(key)
    # 如果用户明确说“详情页/参数/规格”，增加更具体的字段，但仍不凭空写价格 MOQ。
    if any(k in text for k in ["spec", "parameter", "规格", "参数"]):
        for key in ["materials", "colors", "carton"]:
            if key not in requested:
                requested.append(key)
    return requested

def match_category_preset(category: str, language: str) -> Dict[str, Any]:
    key = category.lower().strip()
    presets = CATEGORY_PRESETS_EN if language == "en" else CATEGORY_PRESETS_ZH
    for pattern, preset in presets.items():
        if pattern.lower() in key:
            return preset
    return {}


def generic_category_value(category: str, theme: str, section_key: str, language: str) -> List[str]:
    title = title_case_category(category) if language == "en" else category
    theme_text = title_case_category(theme) if language == "en" else theme
    if language == "en":
        generic = {
            "positioning": [f"{title} presented as a focused catalog category for buyer review."],
            "range": [f"Core {title} direction", "Coordinated variants", "Bundle-ready options"],
            "scenario": [f"Retail display, seasonal programs and {theme_text.lower()} buyer communication"],
            "selling_points": ["Clear visual identity", "Flexible style options", "Easy to expand with real SKU data"],
            "materials": ["Material details should be filled from confirmed product specifications."],
            "colors": ["Color direction should follow the requested visual theme and real product options."],
            "customization": ["Logo, packaging and set configuration can be added after supplier confirmation."],
            "packaging": ["Packaging format should be filled from real product data."],
            "price": ["Price tiers require confirmed quotation data before buyer-facing release."],
            "moq": ["MOQ should be filled only after supplier confirmation."],
            "carton": ["Carton size and packing quantity require confirmed logistics data."],
            "certification": ["Certification details should match the target market and verified documents."],
            "delivery": ["Lead time should be confirmed by production schedule and order quantity."],
            "sku": [f"Representative {title} items", "Color/style variants", "Set or bundle options"],
            "audience": ["Useful for fast category screening before final SKU selection."],
        }
    else:
        generic = {
            "positioning": [f"{title} 作为目录册中的独立类目，用于客户快速判断方向。"],
            "range": [f"{title}基础款", "颜色/规格延展款", "组合装方向"],
            "scenario": [f"零售陈列、季节项目和{theme_text}客户沟通"],
            "selling_points": ["视觉方向清晰", "款式延展灵活", "便于后续补真实 SKU 数据"],
            "materials": ["材质信息需根据真实商品规格补充。"],
            "colors": ["色彩方向应结合用户主题和真实商品可选色。"],
            "customization": ["Logo、包装、组合方式需在供应商确认后补充。"],
            "packaging": ["包装方式需以真实商品资料为准。"],
            "price": ["报价需以确认后的价格表为准，不在初稿中臆造。"],
            "moq": ["MOQ 需供应商确认后填写。"],
            "carton": ["箱规和装箱数量需以物流资料为准。"],
            "certification": ["认证信息需匹配目标市场并附真实证书。"],
            "delivery": ["交期需结合生产排期和订单数量确认。"],
            "sku": [f"{title}代表款", "颜色/规格延展款", "组合套装方向"],
            "audience": ["适合客户快速筛选类目方向，再进入具体 SKU 确认。"],
        }
    return generic.get(section_key, generic["positioning"])


def category_section_items(category: str, theme: str, section_key: str, language: str) -> List[str]:
    preset = match_category_preset(category, language)
    if section_key in preset:
        return list(preset[section_key])
    return generic_category_value(category, theme, section_key, language)


def build_category_sections(category: str, theme: str, language: str, requested_sections: List[str]) -> List[Dict[str, Any]]:
    labels = SECTION_LABELS[language]
    # 没有被用户点名的价格、MOQ、箱规、认证、交期不主动出现，避免目录册看起来像硬模板。
    if requested_sections:
        cleaned_requested = [k for k in requested_sections if k != "positioning"]
        if requested_sections == ["positioning"]:
            section_keys = ["positioning"]
        elif len(cleaned_requested) >= 4:
            # 用户点名的字段优先，不用“产品定位”挤掉用户明确要求的 MOQ/认证/包装等字段。
            section_keys = cleaned_requested
        else:
            section_keys = ["positioning"] + cleaned_requested
    else:
        section_keys = ["positioning", "range"]
        # 对 party/场景型需求，可以加场景；普通商品不强加“适用场景”。
        if any(k in f"{theme} {category}".lower() for k in ["party", "festive", "holiday", "派对", "节庆", "节日"]):
            section_keys.append("scenario")
    result: List[Dict[str, Any]] = []
    seen = set()
    for key in section_keys:
        if key in seen:
            continue
        seen.add(key)
        result.append({"title": labels.get(key, key), "items": category_section_items(category, theme, key, language)})
    return result[:4]


def build_company_intro_sections(prompt: str, brand: str, theme: str, language: str, requested_sections: List[str]) -> Tuple[List[str], List[Dict[str, Any]]]:
    desc = extract_page1_description(prompt)
    if language == "en":
        cleaned = re.sub(r"^Company introduction for\s+[^-–—:：]+\s*[-–—:：]?\s*", "", desc, flags=re.IGNORECASE).strip()
        bullets = [cleaned] if cleaned else [f"Supplier introduction for {title_case_category(theme)} catalog communication."]
        sections = [
            {"title": "Catalog Role", "items": ["Opening page for company positioning and buyer context"]},
            {"title": "Presentation Focus", "items": ["Keep the introduction concise; product pages carry the main category content"]},
        ]
    else:
        cleaned = re.sub(r"^(公司介绍|企业介绍)\s*[:：-]?\s*", "", desc, flags=re.IGNORECASE).strip()
        bullets = [cleaned] if cleaned else [f"用于{theme}目录册沟通的供应商介绍页。"]
        sections = [
            {"title": "页面作用", "items": ["用于说明公司定位和目录背景"]},
            {"title": "表达重点", "items": ["公司介绍保持简洁，主要内容放在后续产品页"]},
        ]
    if "certification" in requested_sections:
        sections.append({"title": SECTION_LABELS[language]["certification"], "items": category_section_items(brand, theme, "certification", language)})
    return bullets[:3], sections[:3]


def choose_slide_variant(style_family: str, index: int, layout: str) -> str:
    if layout in {"company_intro", "cover_catalog"}:
        return layout
    sequences = {
        "minimal": ["image_right", "image_top", "image_left"],
        "luxury": ["image_left", "image_right", "image_top"],
        "playful": ["image_top", "image_left", "image_right"],
        "festive": ["image_left", "image_top", "image_right"],
        "nature": ["image_right", "image_left", "image_top"],
        "tech": ["image_right", "image_top", "image_left"],
        "commercial": ["image_left", "image_right", "image_top"],
    }
    seq = sequences.get(style_family, sequences["commercial"])
    return seq[(index - 1) % len(seq)]

def extract_product_categories(prompt: str, language: str) -> List[str]:
    text = clean_prompt(prompt)
    candidates = ""
    patterns = [
        r"Product\s+categories\s*[:：]\s*(.+?)(?:[\.。]\s*|$)",
        r"产品类目\s*[:：]\s*(.+?)(?:[\.。]\s*|$)",
        r"商品类目\s*[:：]\s*(.+?)(?:[\.。]\s*|$)",
        r"类目\s*[:：]\s*(.+?)(?:[\.。]\s*|$)",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            candidates = match.group(1).strip()
            break
    if not candidates:
        return []

    parts: List[str] = []
    # 先按逗号/分号切，避免把 Balloons and Balloon Sets 里的 and 切掉。
    for item in re.split(r"[,，;；]", candidates):
        item = re.sub(r"^\s*(and|以及|和)\s+", "", item.strip(), flags=re.IGNORECASE)
        item = item.strip(" .。;；")
        if not item:
            continue
        if language == "en":
            item = title_case_category(item)
            if item.lower().startswith("other "):
                item = "Other " + item[6:]
        parts.append(item)
    # 去重但保序。
    seen = set()
    result = []
    for item in parts:
        key = item.lower()
        if key not in seen:
            result.append(item)
            seen.add(key)
    return result


def generated_categories(theme: str, language: str, needed: int) -> List[str]:
    if language == "en":
        base = [
            title_case_category(theme) if theme and theme != "product catalog" else "Hero Products",
            "Bestseller Collection",
            "Premium Gift Sets",
            "Retail Display Options",
            "Packaging and Customization",
            "Buyer Information",
        ]
    else:
        base = [
            f"{theme}主推款" if theme else "主推商品",
            "热销组合",
            "高端礼品款",
            "零售陈列方案",
            "包装与定制",
            "采购沟通信息",
        ]
    while len(base) < needed:
        base.append((f"Catalog Category {len(base)+1}" if language == "en" else f"目录类目 {len(base)+1}"))
    return base[:needed]


def english_category_content(category: str, theme: str) -> List[Dict[str, Any]]:
    # Backward-compatible wrapper. New generation uses build_category_sections().
    return build_category_sections(category, theme, "en", [])


def chinese_category_content(category: str, theme: str) -> List[Dict[str, Any]]:
    # Backward-compatible wrapper. New generation uses build_category_sections().
    return build_category_sections(category, theme, "zh", [])

def make_image_prompt(slide: SlidePlan, plan: DeckPlan) -> str:
    """
    图片生成原则：
    1. 防抄袭、侵权，不等于禁止场景；用户明确要求场景时允许生成场景。
    2. 用户只写 theme/catalog/product category 时，主题只影响色彩与氛围，不能自动扩散成派对房间或活动现场。
    3. 非派对类产品不能因为 theme 里有 party/festive 就被强行放进派对背景。
    4. 防抄袭重点放在原创构图、无品牌、无商标、无 IP、无现成包装、无图库复刻。
    5. 图片里不出现文字，避免乱码。
    """
    scene_mode = os.getenv("PPT_IMAGE_SCENE_MODE", "auto").strip().lower()
    subject = (slide.category or slide.title or plan.theme or "product").strip()

    user_text = " ".join(
        str(x or "")
        for x in [
            getattr(plan, "raw_prompt", ""),
            getattr(plan, "user_prompt", ""),
            getattr(plan, "prompt", ""),
            plan.theme,
            slide.title,
            slide.category,
        ]
    ).lower()

    scene_keywords_en = [
        "scene",
        "lifestyle",
        "in use",
        "use scenario",
        "party scene",
        "birthday party",
        "party setup",
        "room setup",
        "event setup",
        "table setting",
        "decorated room",
        "retail display",
        "showroom",
        "background scene",
        "with background",
    ]
    scene_keywords_zh = [
        "场景",
        "使用场景",
        "生活方式",
        "派对场景",
        "生日派对",
        "派对布置",
        "房间布置",
        "活动现场",
        "陈列场景",
        "展示场景",
        "桌面布置",
        "带背景",
        "有背景",
    ]
    explicit_scene_requested = any(k in user_text for k in scene_keywords_en + scene_keywords_zh)

    def should_use_lifestyle_scene() -> bool:
        if scene_mode in {"product_only", "packshot"}:
            return False
        if scene_mode in {"lifestyle", "scene"}:
            return True
        # auto 模式：只有用户明确要场景时才生成场景。
        return explicit_scene_requested

    use_scene = should_use_lifestyle_scene()

    originality_guard_en = (
        "original generic commercial product design, original composition, "
        "not based on any existing brand, not copied from any stock photo, "
        "no trademark, no copyrighted character, no famous IP, "
        "no recognizable branded packaging, no imitation of existing catalog photos"
    )
    safety_negative_en = (
        "no text, no watermark, no logo, no brand name, no readable label, "
        "no people, no celebrity, no trademark, no copyrighted artwork, "
        "no famous character, no website screenshot"
    )
    originality_guard_zh = (
        "原创通用商业产品设计，原创构图，不基于任何现有品牌，"
        "不复制图库照片，不模仿现有目录册图片，不出现商标、品牌包装、IP角色、版权角色"
    )
    safety_negative_zh = (
        "无文字、无水印、无logo、无品牌名、无可读标签、无人像、无明星、无商标、无版权图案、无网页截图"
    )

    if plan.language == "en":
        if slide.layout == "company_intro":
            if use_scene:
                return (
                    f"Original premium commercial catalog opening scene for {plan.brand_or_company}. "
                    f"A clean, generic, non-branded product display environment with representative products, "
                    f"professional catalog lighting, premium composition, suitable for a supplier introduction page. "
                    f"The scene may reflect the requested theme: {plan.theme}, but must not copy any real catalog or branded display. "
                    f"{originality_guard_en}. {safety_negative_en}."
                )
            return (
                f"Premium commercial catalog opening product arrangement for {plan.brand_or_company}. "
                f"Assorted representative generic non-branded products, clean studio background, "
                f"premium catalog lighting, polished supplier introduction image. "
                f"Use subtle visual accents inspired by {plan.theme}; do not create a full venue unless the user explicitly asks for a scene. "
                f"{originality_guard_en}. {safety_negative_en}."
            )

        if use_scene:
            return (
                f"Original commercial catalog lifestyle scene for {subject}. "
                f"The product must be the clear main subject. "
                f"Create a clean, generic, non-branded scene that matches the user's requested theme: {plan.theme}. "
                f"Use tasteful props and background only when they support the product category. "
                f"Do not let the background overpower the product. "
                f"{originality_guard_en}. {safety_negative_en}."
            )

        return (
            f"Commercial e-commerce catalog product photography of {subject}. "
            f"Isolated product packshot, clean seamless warm-white or light-gray background, "
            f"studio lighting, realistic materials, sharp details, premium catalog composition. "
            f"Use only subtle color accents inspired by {plan.theme}; do not invent a full event venue unless the user explicitly asks for a scene. "
            f"{originality_guard_en}. {safety_negative_en}."
        )

    if slide.layout == "company_intro":
        if use_scene:
            return (
                f"{plan.brand_or_company} 的原创商业目录册开场场景图。"
                f"干净、通用、无品牌的产品展示环境，展示代表性产品，适合供应商介绍页。"
                f"场景可以体现用户要求的主题：{plan.theme}，但不能复制真实目录册、广告或品牌陈列。"
                f"{originality_guard_zh}。{safety_negative_zh}。"
            )
        return (
            f"{plan.brand_or_company} 商业目录册开场产品陈列图，原创通用无品牌产品，"
            f"干净棚拍背景，高级目录册灯光，可以参考“{plan.theme}”的轻微色彩氛围，"
            f"但不要在用户未明确要求时生成完整场景背景。"
            f"{originality_guard_zh}。{safety_negative_zh}。"
        )

    if use_scene:
        return (
            f"{subject} 的原创商业目录册场景图。"
            f"产品必须是主体，场景应符合用户要求的主题：{plan.theme}。"
            f"可以使用适度道具和背景，但不能喧宾夺主，不能模仿现有广告、现有包装或图库构图。"
            f"{originality_guard_zh}。{safety_negative_zh}。"
        )

    return (
        f"{subject} 的商业电商目录册产品图。"
        f"只展示产品本体，干净的暖白色或浅灰色无缝背景，棚拍灯光，真实材质，细节清晰。"
        f"可以参考“{plan.theme}”的轻微色彩氛围，但不要在用户未明确要求时生成完整场景背景。"
        f"{originality_guard_zh}。{safety_negative_zh}。"
    )

def build_deck_plan(prompt: str) -> DeckPlan:
    prompt = clean_prompt(prompt)
    default_page_count = env_int("PPT_DEFAULT_PAGE_COUNT", 5)
    max_page_count = env_int("PPT_MAX_PAGE_COUNT", 20)
    language = detect_language(prompt)
    page_count = extract_page_count(prompt, default_page_count, max_page_count)
    brand = extract_brand_or_company(prompt, language)
    theme = extract_theme(prompt, language)
    style = extract_style(prompt, language)
    style_family = detect_style_family(prompt, theme, language)
    requested_sections = extract_requested_sections(prompt)
    categories = extract_product_categories(prompt, language)

    has_explicit_page_plan = bool(re.search(r"Page\s+1\s*[:：]|Pages?\s+\d+\s*[-–—]\s*\d+", prompt, flags=re.IGNORECASE))
    has_company_intro = bool(re.search(r"company\s+introduction|公司介绍|企业介绍", prompt, flags=re.IGNORECASE))
    has_cover_request = bool(re.search(r"cover\s+page|封面|首页", prompt, flags=re.IGNORECASE))
    mode = "detailed" if (has_explicit_page_plan or categories or requested_sections) else "simple"

    slides: List[SlidePlan] = []

    # 只有用户明确要求公司介绍时才生成公司介绍页；不再因为“英文详细需求”自动塞一页固定公司模板。
    if has_company_intro:
        bullets, sections = build_company_intro_sections(prompt, brand, theme, language, requested_sections)
        subtitle = f"{title_case_category(theme)} Product Catalog" if language == "en" else f"{theme}商品目录册"
        slides.append(
            SlidePlan(
                page=1,
                layout="company_intro",
                title=brand,
                subtitle=subtitle,
                bullets=bullets,
                sections=sections,
                layout_variant="company_intro",
            )
        )
    else:
        # 若用户给的类目数量已经覆盖全部页数，就不额外生成封面。
        # 若未给足类目或明确要封面，则保留一页目录册封面。
        should_add_cover = has_cover_request or not categories or len(categories) < page_count
        if should_add_cover and page_count > 0:
            if language == "en":
                title = title_case_category(theme)
                subtitle = "Product Catalog"
                sections = [
                    {"title": "Catalog Direction", "items": ["A concise catalog draft built from the user's description"]},
                ]
                if requested_sections:
                    sections.append({"title": "Requested Detail Focus", "items": [SECTION_LABELS[language].get(k, k) for k in requested_sections[:4]]})
            else:
                title = theme
                subtitle = "商品目录册"
                sections = [
                    {"title": "目录方向", "items": ["根据用户描述生成的商品目录册初稿"]},
                ]
                if requested_sections:
                    sections.append({"title": "本次重点字段", "items": [SECTION_LABELS[language].get(k, k) for k in requested_sections[:4]]})
            slides.append(SlidePlan(page=1, layout="cover_catalog", title=title, subtitle=subtitle, sections=sections, layout_variant="cover_catalog"))

    remaining = page_count - len(slides)
    if remaining > 0:
        category_plan = categories[:remaining]
        if len(category_plan) < remaining:
            # 只有在页数需要补齐时才生成相关目录页，避免把用户未要求的 MOQ/箱规/认证等字段硬塞进去。
            category_plan.extend(generated_categories(theme, language, remaining - len(category_plan)))
        start_page = len(slides) + 1
        for idx, category in enumerate(category_plan[:remaining], start=start_page):
            sections = build_category_sections(category, theme, language, requested_sections)
            if language == "en":
                subtitle = style_profile_for(style_family)["name_en"]
                bullets = []
            else:
                subtitle = style_profile_for(style_family)["name_zh"]
                bullets = []
            slides.append(
                SlidePlan(
                    page=idx,
                    layout="category_showcase",
                    title=category,
                    subtitle=subtitle,
                    category=category,
                    bullets=bullets,
                    sections=sections,
                    layout_variant=choose_slide_variant(style_family, idx, "category_showcase"),
                )
            )

    slides = slides[:page_count]
    for i, slide in enumerate(slides, start=1):
        slide.page = i
        slide.image_key = f"slide_{i:02d}"
        if not slide.layout_variant:
            slide.layout_variant = choose_slide_variant(style_family, i, slide.layout)

    plan = DeckPlan(
        prompt=prompt,
        language=language,
        page_count=page_count,
        deck_type="product_catalog",
        brand_or_company=brand,
        theme=theme,
        style=style,
        mode=mode,
        slides=slides,
        explicit_categories=categories,
        style_family=style_family,
        requested_sections=requested_sections,
    )
    briefs: List[ImageBrief] = []
    for slide in plan.slides:
        slide.image_prompt = make_image_prompt(slide, plan)
        briefs.append(ImageBrief(key=slide.image_key, page=slide.page, title=slide.title, prompt=slide.image_prompt, aspect_ratio="16:9"))
    plan.image_briefs = briefs
    validate_plan_before_render(plan)
    return plan

def validate_plan_before_render(plan: DeckPlan) -> None:
    if plan.page_count < 1:
        raise RuntimeError("需求解析失败：page_count 不能小于 1。")
    if len(plan.slides) != plan.page_count:
        raise RuntimeError(f"需求解析失败：deck_plan 页数不一致，要求 {plan.page_count} 页，计划 {len(plan.slides)} 页。")
    seen_pages = [s.page for s in plan.slides]
    if seen_pages != list(range(1, plan.page_count + 1)):
        raise RuntimeError(f"需求解析失败：页码不连续：{seen_pages}")
    if plan.language not in {"en", "zh"}:
        raise RuntimeError(f"需求解析失败：不支持的语言：{plan.language}")
    if plan.mode == "detailed" and plan.explicit_categories:
        required = max(0, plan.page_count - 1)
        if len(plan.explicit_categories) < required:
            # 不直接失败：少类目时自动补齐，但写入 metadata。这里做明确提示。
            pass


def plan_to_dict(plan: DeckPlan) -> Dict[str, Any]:
    return {
        **{k: v for k, v in asdict(plan).items() if k not in {"slides", "image_briefs"}},
        "slides": [asdict(s) for s in plan.slides],
        "image_briefs": [asdict(b) for b in plan.image_briefs],
    }


def read_openclaw_minimax_key() -> str:
    config_path = Path(os.getenv("OPENCLAW_CONFIG", "~/.openclaw/openclaw.json")).expanduser()
    if not config_path.exists():
        return ""
    try:
        data = json.loads(config_path.read_text(encoding="utf-8"))
    except Exception:
        return ""

    key_names = {"apikey", "api_key", "key", "token", "authorization"}

    def walk(obj: Any, parent_keys: List[str]) -> Iterable[str]:
        if isinstance(obj, dict):
            for k, v in obj.items():
                lk = str(k).lower()
                new_parent = parent_keys + [lk]
                if isinstance(v, str) and lk in key_names and any("minimax" in p or "minimaxi" in p for p in parent_keys):
                    yield v
                else:
                    yield from walk(v, new_parent)
        elif isinstance(obj, list):
            for item in obj:
                yield from walk(item, parent_keys)

    for candidate in walk(data, []):
        candidate = candidate.strip()
        if candidate and not candidate.startswith("${") and len(candidate) > 8:
            if candidate.lower().startswith("bearer "):
                candidate = candidate.split(None, 1)[1].strip()
            return candidate
    return ""


def get_minimax_api_key() -> str:
    for name in ["MINIMAX_API_KEY", "MINIMAX_TOKEN", "MINIMAX_API_TOKEN"]:
        value = os.getenv(name, "").strip()
        if value:
            return value.removeprefix("Bearer ").strip()
    key_file = os.getenv("MINIMAX_API_KEY_FILE", "").strip()
    if key_file:
        path = Path(key_file).expanduser()
        if path.exists():
            return path.read_text(encoding="utf-8").strip().removeprefix("Bearer ").strip()
    return read_openclaw_minimax_key()


def normalize_minimax_image_api_url(raw_url: str) -> str:
    url = (raw_url or "").strip() or "https://api.minimax.com/v1/image_generation"
    if not url.startswith("http://") and not url.startswith("https://"):
        url = "https://" + url.lstrip("/")
    parsed = urllib.parse.urlparse(url)
    scheme = parsed.scheme or "https"
    host = parsed.netloc
    path = parsed.path or ""
    if not host:
        stripped = url.replace("https://", "").replace("http://", "").strip("/")
        host, _, path_tail = stripped.partition("/")
        path = "/" + path_tail if path_tail else ""
    if "image_generation" in path:
        return urllib.parse.urlunparse((scheme, host, path.rstrip("/"), "", "", ""))
    return urllib.parse.urlunparse((scheme, host, "/v1/image_generation", "", "", ""))


def get_minimax_image_api_candidates() -> List[str]:
    configured_url = (
        os.getenv("MINIMAX_IMAGE_API_URL", "").strip()
        or os.getenv("MINIMAX_IMAGE_BASE_URL", "").strip()
        or os.getenv("MINIMAX_BASE_URL", "").strip()
        or "https://api.minimax.com/v1/image_generation"
    )
    raw_candidates = os.getenv("MINIMAX_IMAGE_API_URL_CANDIDATES", "").strip()
    if raw_candidates:
        candidates = [x.strip() for x in re.split(r"[,\n]", raw_candidates) if x.strip()]
    else:
        candidates = [configured_url, "https://api.minimaxi.com/v1/image_generation", "https://api.minimax.io/v1/image_generation"]
    if not env_bool("MINIMAX_AUTO_FALLBACK_HOSTS", True):
        candidates = [configured_url]
    normalized: List[str] = []
    seen = set()
    for item in candidates:
        url = normalize_minimax_image_api_url(item)
        if url not in seen:
            normalized.append(url)
            seen.add(url)
    return normalized


def assert_hostname_resolves(api_url: str, timeout: int = 8) -> None:
    parsed = urllib.parse.urlparse(api_url)
    host = parsed.hostname
    if not host:
        raise RuntimeError(f"MiniMax 生图接口地址不合法：{api_url}")
    old_timeout = socket.getdefaulttimeout()
    socket.setdefaulttimeout(timeout)
    try:
        socket.getaddrinfo(host, parsed.port or 443)
    except Exception as exc:
        raise RuntimeError(
            f"MiniMax 生图接口域名解析失败：{host}；当前接口地址={api_url}。"
            "如果服务器通过代理访问 MiniMax，建议设置 MINIMAX_NETWORK_PRECHECK=0。"
        ) from exc
    finally:
        socket.setdefaulttimeout(old_timeout)


def collect_strings(obj: Any) -> Iterable[str]:
    if isinstance(obj, str):
        yield obj
    elif isinstance(obj, dict):
        for value in obj.values():
            yield from collect_strings(value)
    elif isinstance(obj, list):
        for value in obj:
            yield from collect_strings(value)


def looks_like_base64_image(text: str) -> bool:
    if text.startswith("data:image/"):
        return True
    if len(text) < 1000:
        return False
    compact = text.strip().replace("\n", "")
    return re.fullmatch(r"[A-Za-z0-9+/=]+", compact) is not None


def decode_base64_image(text: str) -> bytes:
    if text.startswith("data:image/"):
        _, _, text = text.partition(",")
    compact = text.strip().replace("\n", "")
    missing = len(compact) % 4
    if missing:
        compact += "=" * (4 - missing)
    return base64.b64decode(compact, validate=False)


class MiniMaxImageClient:
    def __init__(self) -> None:
        self.api_key = get_minimax_api_key()
        self.api_urls = get_minimax_image_api_candidates()
        self.model = os.getenv("MINIMAX_IMAGE_MODEL", "image-01").strip()
        self.timeout = env_int("MINIMAX_IMAGE_TIMEOUT", 180)
        self.response_format = os.getenv("MINIMAX_IMAGE_RESPONSE_FORMAT", "base64").strip().lower()
        self.prompt_optimizer = env_bool("MINIMAX_PROMPT_OPTIMIZER", False)
        self.verify_ssl = env_bool("MINIMAX_VERIFY_SSL", True)
        self.network_precheck = env_bool("MINIMAX_NETWORK_PRECHECK", False)
        proxy_url = os.getenv("MINIMAX_PROXY_URL", "").strip()
        self.proxies = {"http": proxy_url, "https": proxy_url} if proxy_url else None

    def generate_to_file(self, prompt: str, out_path: Path, aspect_ratio: str = "16:9") -> None:
        if not self.api_key:
            raise RuntimeError("MINIMAX_API_KEY 未配置，且未能从 ~/.openclaw/openclaw.json 自动读取到 MiniMax 密钥；为避免生成纯文字 PPT，本次已停止。")
        errors: List[str] = []
        for api_url in self.api_urls:
            try:
                if self.network_precheck:
                    assert_hostname_resolves(api_url)
                self._generate_to_file_once(api_url, prompt, out_path, aspect_ratio)
                return
            except Exception as exc:
                errors.append(f"{api_url} => {exc}")
                if not env_bool("MINIMAX_AUTO_FALLBACK_HOSTS", True):
                    break
        raise RuntimeError("MiniMax 生图全部候选接口均失败：" + "；".join(errors))

    def _generate_to_file_once(self, api_url: str, prompt: str, out_path: Path, aspect_ratio: str = "16:9") -> None:
        payload: Dict[str, Any] = {
            "model": self.model,
            "prompt": prompt[:1500],
            "aspect_ratio": aspect_ratio,
            "response_format": self.response_format,
            "n": 1,
            "prompt_optimizer": self.prompt_optimizer,
        }
        headers = {"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}
        try:
            resp = requests.post(api_url, headers=headers, json=payload, timeout=self.timeout, verify=self.verify_ssl, proxies=self.proxies)
        except requests.exceptions.RequestException as exc:
            raise RuntimeError(f"MiniMax 生图请求失败：url={api_url}；{exc}") from exc
        if resp.status_code >= 400:
            raise RuntimeError(f"MiniMax 生图失败 HTTP {resp.status_code}: {resp.text[:600]}")
        try:
            data = resp.json()
        except Exception as exc:
            raise RuntimeError(f"MiniMax 生图返回不是 JSON：{exc}; body={resp.text[:300]}") from exc
        base_resp = data.get("base_resp") if isinstance(data, dict) else None
        if isinstance(base_resp, dict) and int(base_resp.get("status_code", 0) or 0) != 0:
            raise RuntimeError(f"MiniMax 生图失败：{base_resp.get('status_msg') or base_resp}")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        image_urls = []
        if isinstance(data, dict) and isinstance(data.get("data"), dict):
            image_urls = data.get("data", {}).get("image_urls", [])
        if image_urls:
            self._download_image(str(image_urls[0]), out_path)
            return
        for item in collect_strings(data):
            if item.startswith("http://") or item.startswith("https://"):
                self._download_image(item, out_path)
                return
            if looks_like_base64_image(item):
                raw = decode_base64_image(item)
                if len(raw) > 1000:
                    out_path.write_bytes(raw)
                    self._normalize_image(out_path)
                    return
        raise RuntimeError(f"MiniMax 生图响应中未找到 image_urls 或 base64 图片字段：{json.dumps(data, ensure_ascii=False)[:800]}")

    def _download_image(self, url: str, out_path: Path) -> None:
        resp = requests.get(url, timeout=self.timeout, verify=self.verify_ssl, proxies=self.proxies)
        if resp.status_code >= 400:
            raise RuntimeError(f"MiniMax 图片下载失败 HTTP {resp.status_code}: {url}")
        out_path.write_bytes(resp.content)
        self._normalize_image(out_path)

    def _normalize_image(self, path: Path) -> None:
        try:
            with Image.open(path) as im:
                im = im.convert("RGB")
                im.save(path, format="PNG")
        except Exception as exc:
            raise RuntimeError(f"MiniMax 图片文件无法被 Pillow 读取：{path}；{exc}") from exc


def find_font() -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for path in candidates:
        try:
            return ImageFont.truetype(path, 36)
        except Exception:
            continue
    return ImageFont.load_default()


def generate_placeholder_image(title: str, out_path: Path, aspect_ratio: str, language: str = "en") -> None:
    size = (1600, 900) if aspect_ratio == "16:9" else (1200, 1200)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGB", size, (248, 250, 252))
    draw = ImageDraw.Draw(im)
    draw.rounded_rectangle([46, 46, size[0] - 46, size[1] - 46], radius=42, fill=(255, 255, 255), outline=(203, 213, 225), width=4)
    draw.ellipse([size[0] * 0.08, size[1] * 0.16, size[0] * 0.34, size[1] * 0.62], fill=(219, 234, 254))
    draw.rounded_rectangle([size[0] * 0.40, size[1] * 0.20, size[0] * 0.80, size[1] * 0.68], radius=34, fill=(224, 242, 254))
    draw.line([size[0] * 0.14, size[1] * 0.74, size[0] * 0.84, size[1] * 0.74], fill=(37, 99, 235), width=9)
    font = find_font()
    label = "OFFLINE TEST IMAGE" if language == "en" else "离线测试图"
    text = f"{label}\n{title[:54]}"
    draw.multiline_text((80, size[1] - 178), text, fill=(71, 85, 105), font=font, spacing=10)
    im.save(out_path, format="PNG")


def prepare_images(plan: DeckPlan, run_dir: Path, image_mode: str) -> List[ImageAsset]:
    image_mode = image_mode.lower().strip()
    if image_mode not in {"minimax", "placeholder", "off"}:
        raise ValueError("PPT_IMAGE_MODE 只能是 minimax / placeholder / off")
    if image_mode == "off":
        if env_bool("PPT_REQUIRE_IMAGES", True):
            raise RuntimeError("PPT_IMAGE_MODE=off 但 PPT_REQUIRE_IMAGES=1；为避免纯文字 PPT，本次已停止。")
        return []
    asset_dir = run_dir / "assets" / image_mode
    client = MiniMaxImageClient() if image_mode == "minimax" else None
    assets: List[ImageAsset] = []
    for brief in plan.image_briefs:
        out_path = asset_dir / f"{brief.key}.png"
        if image_mode == "minimax":
            assert client is not None
            client.generate_to_file(brief.prompt, out_path, brief.aspect_ratio)
        else:
            generate_placeholder_image(brief.title, out_path, brief.aspect_ratio, plan.language)
        assets.append(ImageAsset(key=brief.key, title=brief.title, prompt=brief.prompt, path=str(out_path), source=image_mode, aspect_ratio=brief.aspect_ratio, page=brief.page))
    if env_bool("PPT_REQUIRE_IMAGES", True) and len(assets) != len(plan.slides):
        raise RuntimeError(f"图片数量校验失败：计划 {len(plan.slides)} 张，实际 {len(assets)} 张。")
    return assets


def asset_map(assets: List[ImageAsset]) -> Dict[str, ImageAsset]:
    return {asset.key: asset for asset in assets}


def set_slide_background(slide, rgb=(248, 250, 252)) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def font_name(language: str) -> str:
    return "Arial" if language == "en" else "Microsoft YaHei"


def add_textbox(slide, left, top, width, height, text, font_size=18, bold=False, color=(31, 41, 55), align=None, language="en"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.clear()
    lines = str(text).split("\n") if str(text) else [""]
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*color)
            run.font.name = font_name(language)
    return box


def add_bullet_box(slide, left, top, width, height, lines: List[str], font_size=10.5, color=(71, 85, 105), language="en"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(0)
    tf.margin_right = Pt(0)
    tf.margin_top = Pt(0)
    tf.margin_bottom = Pt(0)
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.size = Pt(font_size)
        p.font.name = font_name(language)
        p.font.color.rgb = RGBColor(*color)
        p.space_after = Pt(3)
    return box


def add_section_card(
    slide,
    left,
    top,
    width,
    height,
    title: str,
    items: List[str],
    accent=(37, 99, 235),
    language="en",
    panel=(255, 255, 255),
    text_color=(15, 23, 42),
    muted=(71, 85, 105),
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(*panel)
    shape.line.color.rgb = RGBColor(226, 232, 240)
    shape.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.07), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*accent)
    bar.line.fill.background()
    add_textbox(slide, left + Inches(0.22), top + Inches(0.13), width - Inches(0.35), Inches(0.3), title, 12.4, True, text_color, language=language)
    add_bullet_box(slide, left + Inches(0.22), top + Inches(0.52), width - Inches(0.42), height - Inches(0.58), items, 9.2 if language == "en" else 9.0, muted, language=language)


def add_image(slide, asset: Optional[ImageAsset], left, top, width, height, label: str, language="en") -> None:
    if asset and Path(asset.path).exists():
        try:
            pic = slide.shapes.add_picture(asset.path, left, top, width=width, height=height)
            pic.line.color.rgb = RGBColor(226, 232, 240)
            return
        except Exception:
            pass
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(226, 232, 240)
    shape.line.color.rgb = RGBColor(203, 213, 225)
    fallback = "Image unavailable" if language == "en" else "图片不可用"
    add_textbox(slide, left + Inches(0.15), top + height / 2 - Inches(0.16), width - Inches(0.3), Inches(0.4), fallback, 10.5, True, (100, 116, 139), PP_ALIGN.CENTER, language=language)


def add_style_accent(slide, plan: DeckPlan) -> None:
    profile = style_profile_for(plan)
    accent = profile["accent"]
    accent2 = profile["accent2"]
    # 只用轻量装饰，避免喧宾夺主；不同 style_family 呈现不同气质。
    if plan.style_family in {"playful", "festive"}:
        for x, y, size, color in [
            (11.55, 0.22, 0.62, accent2),
            (12.25, 0.82, 0.34, accent),
            (0.22, 6.42, 0.46, profile["accent3"]),
        ]:
            shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor(*color)
            shp.line.fill.background()
    elif plan.style_family in {"minimal", "luxury", "nature"}:
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.64), Inches(0.27), Inches(0.08), Inches(6.68))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(*accent)
        line.line.fill.background()
    elif plan.style_family == "tech":
        for y in [0.42, 6.82]:
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(y), Inches(11.95), Inches(0.018))
            line.fill.solid()
            line.fill.fore_color.rgb = RGBColor(*accent2)
            line.line.fill.background()


def add_footer(slide, plan: DeckPlan, page: int) -> None:
    if plan.language == "en":
        text = f"{plan.brand_or_company} · {title_case_category(plan.theme)} · {page}/{plan.page_count}"
    else:
        text = f"{plan.brand_or_company} · {plan.theme} · {page}/{plan.page_count}"
    add_textbox(slide, Inches(0.66), Inches(7.08), Inches(12.0), Inches(0.22), text, 8.2, False, (148, 163, 184), language=plan.language)


def section_positions_for(variant: str, count: int) -> List[Tuple[Any, Any, Any, Any]]:
    count = max(1, min(count, 4))
    if variant == "image_right":
        if count == 1:
            return [(Inches(0.9), Inches(2.05), Inches(4.95), Inches(1.85))]
        if count == 2:
            return [(Inches(0.9), Inches(2.0), Inches(4.95), Inches(1.65)), (Inches(0.9), Inches(3.92), Inches(4.95), Inches(1.65))]
        return [(Inches(0.9), Inches(1.75 + 1.46 * i), Inches(4.95), Inches(1.22)) for i in range(count)]
    if variant == "image_top":
        if count == 1:
            return [(Inches(1.0), Inches(5.08), Inches(11.1), Inches(1.08))]
        if count == 2:
            return [(Inches(1.0 + 5.65 * i), Inches(5.02), Inches(5.35), Inches(1.18)) for i in range(2)]
        return [(Inches(1.0 + (i % 3) * 3.78), Inches(5.0), Inches(3.48), Inches(1.18)) for i in range(min(count, 3))]
    # image_left default
    if count == 1:
        return [(Inches(6.82), Inches(2.05), Inches(5.35), Inches(1.85))]
    if count == 2:
        return [(Inches(6.82), Inches(1.85), Inches(5.35), Inches(1.55)), (Inches(6.82), Inches(3.72), Inches(5.35), Inches(1.55))]
    return [(Inches(6.78), Inches(1.5 + 1.46 * i), Inches(5.55), Inches(1.22)) for i in range(count)]


def create_company_intro_slide(prs: Presentation, plan: DeckPlan, slide_plan: SlidePlan, assets: Dict[str, ImageAsset]) -> None:
    profile = style_profile_for(plan)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, profile["bg"])
    add_style_accent(slide, plan)
    add_textbox(slide, Inches(0.92), Inches(0.62), Inches(5.65), Inches(0.78), slide_plan.title, 34, True, profile["text"], language=plan.language)
    add_textbox(slide, Inches(0.95), Inches(1.42), Inches(5.55), Inches(0.36), slide_plan.subtitle, 15.5, False, profile["accent"], language=plan.language)
    if slide_plan.bullets:
        add_bullet_box(slide, Inches(0.98), Inches(2.02), Inches(5.25), Inches(0.92), slide_plan.bullets, 10.2, profile["muted"], plan.language)

    positions = [
        (Inches(0.95), Inches(3.18), Inches(2.62), Inches(1.5)),
        (Inches(3.78), Inches(3.18), Inches(2.62), Inches(1.5)),
        (Inches(0.95), Inches(5.02), Inches(5.45), Inches(1.18)),
    ]
    accents = [profile["accent"], profile["accent2"], profile["accent3"]]
    for idx, section in enumerate(slide_plan.sections[:3]):
        left, top, width, height = positions[idx]
        add_section_card(slide, left, top, width, height, section["title"], section["items"], accents[idx], plan.language, profile["panel"], profile["text"], profile["muted"])

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.68), Inches(0.72), Inches(5.9), Inches(5.9))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(*profile["panel"])
    panel.line.color.rgb = RGBColor(226, 232, 240)
    add_image(slide, assets.get(slide_plan.image_key), Inches(6.92), Inches(0.98), Inches(5.42), Inches(4.98), slide_plan.title, plan.language)
    add_footer(slide, plan, slide_plan.page)


def create_cover_slide(prs: Presentation, plan: DeckPlan, slide_plan: SlidePlan, assets: Dict[str, ImageAsset]) -> None:
    profile = style_profile_for(plan)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, profile["bg"])
    add_style_accent(slide, plan)
    title_size = 34 if plan.language == "en" else 32
    add_textbox(slide, Inches(0.95), Inches(0.84), Inches(5.6), Inches(0.88), slide_plan.title, title_size, True, profile["text"], language=plan.language)
    add_textbox(slide, Inches(0.98), Inches(1.76), Inches(5.25), Inches(0.36), slide_plan.subtitle, 15.5, False, profile["accent"], language=plan.language)
    for idx, section in enumerate(slide_plan.sections[:2]):
        add_section_card(slide, Inches(0.98), Inches(2.62 + 1.48 * idx), Inches(5.25), Inches(1.16), section["title"], section["items"], [profile["accent"], profile["accent2"]][idx], plan.language, profile["panel"], profile["text"], profile["muted"])
    add_image(slide, assets.get(slide_plan.image_key), Inches(6.58), Inches(0.72), Inches(5.95), Inches(5.9), slide_plan.title, plan.language)
    add_footer(slide, plan, slide_plan.page)


def create_category_showcase_slide(prs: Presentation, plan: DeckPlan, slide_plan: SlidePlan, assets: Dict[str, ImageAsset]) -> None:
    profile = style_profile_for(plan)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, profile["bg"])
    add_style_accent(slide, plan)
    variant = slide_plan.layout_variant or choose_slide_variant(plan.style_family, slide_plan.page, "category_showcase")
    title_left = Inches(0.92 if variant != "image_left" else 0.72)
    add_textbox(slide, title_left, Inches(0.48), Inches(8.8), Inches(0.52), slide_plan.title, 26.5, True, profile["text"], language=plan.language)
    add_textbox(slide, title_left, Inches(1.02), Inches(8.2), Inches(0.28), slide_plan.subtitle, 10.8, False, profile["muted"], language=plan.language)

    asset = assets.get(slide_plan.image_key)
    if variant == "image_right":
        add_image(slide, asset, Inches(6.48), Inches(1.48), Inches(5.72), Inches(4.92), slide_plan.title, plan.language)
    elif variant == "image_top":
        add_image(slide, asset, Inches(1.0), Inches(1.45), Inches(11.15), Inches(3.28), slide_plan.title, plan.language)
    else:
        add_image(slide, asset, Inches(0.8), Inches(1.48), Inches(5.6), Inches(4.92), slide_plan.title, plan.language)

    accents = [profile["accent"], profile["accent2"], profile["accent3"], (245, 158, 11)]
    sections = slide_plan.sections[:4]
    for idx, section in enumerate(sections):
        left, top, width, height = section_positions_for(variant, len(sections))[idx]
        add_section_card(slide, left, top, width, height, section["title"], section["items"], accents[idx], plan.language, profile["panel"], profile["text"], profile["muted"])
    add_footer(slide, plan, slide_plan.page)

def build_presentation(plan: DeckPlan, pptx_path: Path, assets: List[ImageAsset]) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    by_key = asset_map(assets)
    for slide_plan in plan.slides:
        if slide_plan.layout == "company_intro":
            create_company_intro_slide(prs, plan, slide_plan, by_key)
        elif slide_plan.layout == "cover_catalog":
            create_cover_slide(prs, plan, slide_plan, by_key)
        elif slide_plan.layout == "category_showcase":
            create_category_showcase_slide(prs, plan, slide_plan, by_key)
        else:
            raise RuntimeError(f"未知版式：{slide_plan.layout}")
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))


def workspace_root() -> Path:
    return Path(__file__).resolve().parents[1]


def choose_render_backend() -> str:
    backend = os.getenv("PPT_RENDER_BACKEND", "auto").strip().lower()
    if backend not in {"auto", "pptmaster", "python_pptx"}:
        raise RuntimeError("PPT_RENDER_BACKEND 只能是 auto / pptmaster / python_pptx")
    return backend


def render_presentation(plan: DeckPlan, pptx_path: Path, assets: List[ImageAsset], run_dir: Path) -> Dict[str, Any]:
    """渲染 PPT。默认 auto：优先 PPT-master，缺失或失败时降级 python-pptx。"""
    backend = choose_render_backend()
    strict = env_bool("PPT_MASTER_STRICT", backend == "pptmaster")

    if backend in {"auto", "pptmaster"}:
        if build_with_pptmaster is None:
            if strict:
                raise RuntimeError("PPT-master 适配器加载失败，且当前要求严格使用 PPT-master。")
        else:
            try:
                info = build_with_pptmaster(plan, pptx_path, assets, run_dir, workspace_root())
                info["backend"] = "pptmaster"
                return info
            except (PPTMasterUnavailable, PPTMasterPipelineError, RuntimeError) as exc:
                if strict:
                    raise RuntimeError(f"PPT-master 渲染失败：{exc}") from exc
                warn_path = run_dir / "PPT_MASTER_FALLBACK.txt"
                warn_path.write_text(
                    "PPT-master 渲染不可用，已按 auto 策略降级到 python-pptx。\n"
                    f"原因：{exc}\n"
                    "如需强制失败，请设置 PPT_RENDER_BACKEND=pptmaster 或 PPT_MASTER_STRICT=1。\n",
                    encoding="utf-8",
                )

    build_presentation(plan, pptx_path, assets)
    return {
        "backend": "python_pptx",
        "fallback_reason": "PPT_RENDER_BACKEND=python_pptx or PPT-master unavailable in auto mode",
    }


def extract_pptx_texts(pptx_path: Path) -> List[str]:
    prs = Presentation(str(pptx_path))
    texts: List[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(str(shape.text))
            if hasattr(shape, "table"):
                try:
                    for row in shape.table.rows:
                        for cell in row.cells:
                            if cell.text:
                                texts.append(cell.text)
                except Exception:
                    pass
    return texts


def normalize_match_text(text: str) -> str:
    """Normalize user/category/PPT text for robust category matching.

    PPT-master conversion can split text into runs/lines or replace '&' with
    natural language. Category validation must not fail because a long title is
    wrapped, hyphenated, or extracted with extra whitespace.
    """
    value = (text or "").lower().replace("&", " and ")
    value = re.sub(r"[^0-9a-z\u3400-\u4dbf\u4e00-\u9fff]+", "", value)
    return value


def category_match_tokens(text: str) -> List[str]:
    raw = (text or "").lower().replace("&", " and ")
    words = re.findall(r"[0-9a-z\u3400-\u4dbf\u4e00-\u9fff]+", raw)
    stop = {"and", "or", "of", "for", "with", "the", "a", "an", "other"}
    return [w for w in words if w not in stop and len(w) > 1]


def category_present(category: str, candidates: Iterable[str]) -> bool:
    expected = normalize_match_text(category)
    if not expected:
        return True
    normalized_candidates = [normalize_match_text(c) for c in candidates if c]
    if any(expected in c or c in expected for c in normalized_candidates if c):
        return True

    expected_tokens = category_match_tokens(category)
    if not expected_tokens:
        return False
    for cand in candidates:
        cand_tokens = set(category_match_tokens(cand))
        if cand_tokens and all(token in cand_tokens for token in expected_tokens):
            return True
    return False


def expected_explicit_categories_for_plan(plan: DeckPlan) -> List[str]:
    """Return explicit categories that are expected to appear as category pages.

    When page 1 is company intro/cover, capacity is the number of actual
    category slides, not simply page_count - 1. This avoids false failures for
    decks that intentionally use all pages as category pages or use a mixed
    structure.
    """
    category_slide_count = sum(1 for s in plan.slides if s.category or s.layout == "category_showcase")
    return plan.explicit_categories[:category_slide_count]


def validate_output(plan: DeckPlan, pptx_path: Path, assets: List[ImageAsset]) -> ValidationResult:
    result = ValidationResult(ok=True)
    errors: List[str] = []
    checks: List[str] = []
    if not pptx_path.exists():
        errors.append(f"PPT 文件不存在：{pptx_path}")
    else:
        prs = Presentation(str(pptx_path))
        actual = len(prs.slides)
        if actual != plan.page_count:
            errors.append(f"页数校验失败：用户/计划要求 {plan.page_count} 页，实际生成 {actual} 页。")
        else:
            checks.append(f"页数校验通过：{actual}/{plan.page_count}")

    if env_bool("PPT_REQUIRE_IMAGES", True):
        if len(assets) != len(plan.slides):
            errors.append(f"图片数量校验失败：计划 {len(plan.slides)} 张，实际 {len(assets)} 张。")
        elif any(not Path(a.path).exists() for a in assets):
            errors.append("图片文件校验失败：存在缺失图片。")
        else:
            checks.append(f"图片数量校验通过：{len(assets)} 张")

    if pptx_path.exists():
        extracted_texts = extract_pptx_texts(pptx_path)
        text_blob = "\n".join(extracted_texts)

        # 类目完整性以 deck_plan 为硬校验；PPT 文本抽取只作为辅助校验。
        # PPT-master 的 SVG→PPTX 过程可能会把长标题拆成多个 text run/换行，
        # 不能因为 python-pptx 抽取文本不完全就把已生成文件判失败。
        expected_categories = expected_explicit_categories_for_plan(plan)
        planned_category_titles = [s.category or s.title for s in plan.slides if s.category or s.layout == "category_showcase"]
        missing_from_plan = [c for c in expected_categories if not category_present(c, planned_category_titles)]
        if missing_from_plan:
            errors.append("用户指定类目未进入生成计划：" + ", ".join(missing_from_plan))
        elif expected_categories:
            checks.append("用户指定类目计划校验通过")

        missing_from_extracted_text = [c for c in expected_categories if not category_present(c, extracted_texts + [text_blob])]
        if missing_from_extracted_text:
            checks.append("PPT 文本抽取未完整命中部分长类目，已按 deck_plan 放行：" + ", ".join(missing_from_extracted_text))
        elif expected_categories:
            checks.append("用户指定类目文本校验通过")
        if plan.language == "en":
            if CJK_RE.search(text_blob):
                errors.append("语言校验失败：用户要求英文，但 PPT 页面文本中仍包含中文字符。")
            else:
                checks.append("英文语言校验通过：页面文本未发现中文字符")
        for idx, slide in enumerate(plan.slides, start=1):
            if not slide.title:
                errors.append(f"第 {idx} 页标题为空。")
        if all(s.title for s in plan.slides):
            checks.append("标题完整性校验通过")

    result.errors = errors
    result.checks = checks
    result.ok = not errors
    return result


def sanitize_filename(text: str, max_chars: int = 56) -> str:
    text = clean_prompt(text)
    for ch in INVALID_FILENAME_CHARS:
        text = text.replace(ch, "")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[，,。.!！?？：:；;]+", "-", text)
    text = re.sub(r"-+", "-", text).strip("-._ ")
    if not text:
        text = "product-catalog"
    if len(text) > max_chars:
        text = text[:max_chars].rstrip("-._ ")
    return text


def make_run_paths(prompt: str, root: Path, now: datetime) -> Dict[str, Path | str]:
    max_chars = env_int("PPT_FILENAME_MAX_CHARS", 56)
    time_prefix = now.strftime("%Y-%m-%d-%H时%M分")
    slug = sanitize_filename(prompt, max_chars=max_chars)
    base_name = f"{time_prefix}-{slug}"
    run_name = base_name
    run_dir = root / run_name
    if env_bool("PPT_AUTO_DEDUP", True):
        i = 1
        while run_dir.exists():
            run_name = f"{base_name}-{i:02d}"
            run_dir = root / run_name
            i += 1
    pptx_path = run_dir / f"{run_name}.pptx"
    return {"run_name": run_name, "run_dir": run_dir, "pptx_path": pptx_path}


def windows_path_for(pptx_path: Path, root: Path) -> str:
    prefix = os.getenv("PPT_WINDOWS_SHARE_PREFIX", "Z:\\yaq\\ppt\\catalog").rstrip("\\/")
    try:
        rel = pptx_path.relative_to(root)
    except ValueError:
        rel = pptx_path.name
    return f"{prefix}\\{str(rel).replace('/', '\\')}"


def write_run_files(run_dir: Path, plan: DeckPlan, result: GenerationResult, now: datetime) -> None:
    run_dir.mkdir(parents=True, exist_ok=True)
    plan_path = run_dir / "deck_plan.json"
    plan_path.write_text(json.dumps(plan_to_dict(plan), ensure_ascii=False, indent=2), encoding="utf-8")
    metadata = {
        "created_at": now.isoformat(timespec="seconds"),
        "timezone": os.getenv("PPT_TIMEZONE", DEFAULT_TZ),
        "output_rule": "成功时飞书只回复 windows_path；失败时回复明确错误。",
        "quality_rule": "deck_plan 是唯一事实源；渲染后校验页数、语言、类目、标题、图片数量。",
        "ppt_master_rule": "默认优先使用 workspace/vendor/ppt-master；安装脚本不会删除或覆盖 vendor/ppt-master。",
        "result": asdict(result),
    }
    (run_dir / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "README.txt").write_text(
        f"OpenClaw PPT 商品目录册生成结果\n"
        f"用户问题：{plan.prompt}\n"
        f"页数：{plan.page_count}\n"
        f"语言：{plan.language}\n"
        f"PPT：{result.windows_path}\n"
        f"deck_plan：{plan_path}\n"
        f"渲染后端：{result.render_backend}\n"
        f"PPT-master 项目：{result.pptmaster_project_dir}\n"
        f"飞书回复规则：成功只回复 PPT 本地映射路径；失败回复错误。\n",
        encoding="utf-8",
    )


def generate(prompt: str, sender_name: str = "", sender_open_id: str = "", root: Optional[str] = None, image_mode: Optional[str] = None) -> GenerationResult:
    start = time.time()
    if load_dotenv:
        load_dotenv()
    now = now_in_configured_timezone()
    prompt = clean_prompt(prompt)
    root_path = Path(root or os.getenv("PPT_CATALOG_ROOT", "/data/share/yaq/ppt/catalog")).expanduser().resolve()
    paths = make_run_paths(prompt, root_path, now)
    run_name = str(paths["run_name"])
    run_dir = Path(paths["run_dir"])
    pptx_path = Path(paths["pptx_path"])
    windows_path = windows_path_for(pptx_path, root_path)
    selected_image_mode = (image_mode or os.getenv("PPT_IMAGE_MODE", "minimax")).strip().lower()

    try:
        plan = build_deck_plan(prompt)
        plan_path = run_dir / "deck_plan.json"
        run_dir.mkdir(parents=True, exist_ok=True)
        plan_path.write_text(json.dumps(plan_to_dict(plan), ensure_ascii=False, indent=2), encoding="utf-8")
        image_assets = prepare_images(plan, run_dir, selected_image_mode)
        render_info = render_presentation(plan, pptx_path, image_assets, run_dir)
        validation = validate_output(plan, pptx_path, image_assets)
        if not validation.ok:
            raise RuntimeError("；".join(validation.errors))
        elapsed = round(time.time() - start, 1)
        result = GenerationResult(
            ok=True,
            prompt=prompt,
            page_count=plan.page_count,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path,
            elapsed_seconds=elapsed,
            language=plan.language,
            deck_plan_path=str(plan_path),
            image_mode=selected_image_mode,
            render_backend=str(render_info.get("backend", "")),
            pptmaster_project_dir=str(render_info.get("project_dir", "")),
            pptmaster_log_path=str(render_info.get("log_path", "")),
            image_count=len(image_assets),
            image_assets=[asdict(asset) for asset in image_assets],
            validation=asdict(validation),
            reply_text=windows_path,
        )
        write_run_files(run_dir, plan, result, now)
        return result
    except Exception as exc:
        elapsed = round(time.time() - start, 1)
        try:
            run_dir.mkdir(parents=True, exist_ok=True)
            (run_dir / "ERROR.txt").write_text(str(exc), encoding="utf-8")
        except Exception:
            pass
        return GenerationResult(
            ok=False,
            prompt=prompt,
            page_count=0,
            run_name=run_name,
            run_dir=str(run_dir),
            pptx_path=str(pptx_path),
            windows_path=windows_path,
            elapsed_seconds=elapsed,
            image_mode=selected_image_mode,
            render_backend=choose_render_backend() if os.getenv("PPT_RENDER_BACKEND", "auto") else "auto",
            error=str(exc),
        )


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate a product catalog PPT from a Feishu/OpenClaw prompt")
    parser.add_argument("--prompt", required=True, help="飞书用户原始问题")
    parser.add_argument("--sender-name", default="", help="飞书发送人姓名，仅写入元数据")
    parser.add_argument("--sender-open-id", default="", help="飞书发送人 open_id，仅写入元数据")
    parser.add_argument("--root", default=None, help="输出根目录，不填则读 PPT_CATALOG_ROOT")
    parser.add_argument("--image-mode", choices=["minimax", "placeholder", "off"], default=None, help="生产用 minimax；离线布局测试用 placeholder")
    parser.add_argument("--json", action="store_true", help="以 JSON 输出")
    args = parser.parse_args()

    result = generate(args.prompt, args.sender_name, args.sender_open_id, args.root, args.image_mode)
    data = asdict(result)
    if args.json:
        print(json.dumps(data, ensure_ascii=False, indent=2))
    else:
        if result.ok:
            print(result.windows_path)
        else:
            print(f"PPT 生成失败：{result.error}", file=sys.stderr)
            sys.exit(1)


if __name__ == "__main__":
    main()
